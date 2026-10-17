# -*- coding: utf-8 -*-
"""Genera el documento de Requerimientos Funcionales corregidos en .docx."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

RUTA = r"C:\Users\RYZEN\Documents\TESIS 2026_DOCUMENTOS\TESIS 2026 ALIANY\Requerimientos_Funcionales.docx"

GRIS = RGBColor(0x55, 0x55, 0x55)

bloques = [
    ("Gestión de clientes", [
        ("RF-01", "Registro de clientes", "Permite registrar nuevos clientes, personas físicas o jurídicas, ingresando datos personales/comerciales y financieros necesarios para la evaluación (nombre o razón social, tipo y número de documento CI/RUC, edad, ingresos, historial de pagos, nivel de endeudamiento y demás parámetros).", "Alta"),
        ("RF-02", "Gestión de clientes", "Permite visualizar, editar y eliminar (baja lógica) los registros de clientes almacenados, de acuerdo con los permisos del usuario.", "Alta"),
        ("RF-03", "Validación de datos del cliente", "Valida los datos antes de almacenarlos: unicidad del número de documento, mayoría de edad, formato válido del teléfono y coherencia de la fecha de nacimiento.", "Alta"),
        ("RF-04", "Gestión de ubicaciones geográficas", "Permite administrar y consultar ubicaciones (país, departamento y ciudad) para asociarlas al domicilio del cliente durante su registro.", "Media"),
    ]),
    ("Motor de reglas y scoring", [
        ("RF-05", "Configuración del motor de reglas", "Permite a los administradores crear y configurar motores de reglas IF-THEN, definiendo parámetros, operadores lógicos, pesos y reglas determinantes.", "Alta"),
        ("RF-06", "Configuración de modelos, factores y umbrales", "Permite configurar modelos de scoring, sus factores por categoría y los umbrales de clasificación (bajo, medio y alto).", "Alta"),
        ("RF-07", "Evaluación de scoring mediante reglas configurables", "Calcula automáticamente el puntaje de riesgo mediante un motor de reglas IF-THEN configurable, procesando parámetros con operadores lógicos, pesos y reglas determinantes, usando los datos y el historial crediticio del cliente.", "Alta"),
        ("RF-08", "Clasificación de riesgo", "Clasifica automáticamente en bajo, medio, alto o rechazado según el puntaje. \u201cRechazado\u201d se asigna al incumplir una regla determinante o al detectar coincidencia en listas de sanciones.", "Alta"),
        ("RF-09", "Verificación contra listas de sanciones (ONU/OFAC)", "Verifica al cliente contra un repositorio local de listas de sanciones (ONU/OFAC) cotejando nombre, apellido y documento por coincidencia parcial; una coincidencia confirmada determina riesgo alto/rechazado.", "Alta"),
        ("RF-10", "Gestión del historial crediticio", "Permite registrar y actualizar el historial crediticio (atrasos, deuda total, endeudamiento, lista negra, antigüedad y tipo de empleo), usado como insumo de la evaluación.", "Alta"),
        ("RF-11", "Generación de explicación y recomendación", "Genera automáticamente una explicación del resultado y una recomendación asociada, según la clasificación de riesgo y las reglas aplicadas.", "Media"),
        ("RF-12", "Reutilización de modelos de scoring", "Permite clonar o reciclar un modelo de scoring existente hacia otra empresa, copiando factores, reglas y umbrales.", "Media"),
        ("RF-13", "Evaluación por criterios ponderados", "Permite evaluar mediante criterios ponderados configurables (capacidad de pago, historial, estabilidad laboral, documentación, patrimonio), calculando un puntaje global y un resultado (aprobado, observado o rechazado).", "Media"),
    ]),
    ("Evaluaciones, reportes y estadísticas", [
        ("RF-14", "Consulta y filtrado de evaluaciones", "Permite consultar y filtrar evaluaciones por fecha, cliente y clasificación de riesgo, según el alcance de empresa del usuario.", "Alta"),
        ("RF-15", "Generación de reportes", "Genera reportes de las evaluaciones en pantalla y en PDF, con datos del cliente, puntaje, clasificación y resultados.", "Alta"),
        ("RF-16", "Tablero de estadísticas (dashboard)", "Presenta un tablero con estadísticas agregadas: total de evaluaciones, distribución por categoría de riesgo y total de clientes.", "Media"),
    ]),
    ("Seguridad, usuarios y multi-empresa", [
        ("RF-17", "Autenticación de usuarios con doble factor", "Permite el inicio de sesión con usuario y contraseña, reforzado con un segundo factor (2FA/TOTP) mediante aplicación de autenticación.", "Alta"),
        ("RF-18", "Configuración y restablecimiento de 2FA", "Permite configurar el 2FA en el primer acceso mediante código QR y permite a un administrador restablecerlo cuando sea necesario.", "Alta"),
        ("RF-19", "Gestión de contraseñas", "Permite cambiar la contraseña y recuperarla mediante código OTP, almacenándolas de forma segura mediante hash.", "Alta"),
        ("RF-20", "Control de acceso por roles", "Restringe el acceso a las funcionalidades según el rol y permisos de cada usuario (propietario, administrador y operador/comercial).", "Alta"),
        ("RF-21", "Gestión de usuarios", "Permite a los administradores dar de alta, editar y desactivar (baja lógica) usuarios y asignarles su rol.", "Alta"),
        ("RF-22", "Gestión multi-empresa (multi-tenant)", "Permite administrar múltiples empresas de forma aislada, garantizando que cada una acceda solo a sus clientes, evaluaciones y configuraciones.", "Media"),
        ("RF-23", "Auditoría de operaciones y evaluaciones", "Registra una bitácora de operaciones y evaluaciones (usuario, módulo, entidad afectada, empresa, fecha/hora y resultado: éxito o fallo) y permite su consulta y filtrado.", "Alta"),
    ]),
    ("Modelo comercial: planes, facturación y pagos", [
        ("RF-24", "Gestión de planes de facturación", "Permite al propietario crear y administrar planes de facturación (cantidad de reportes incluidos, precio del plan, precio de referencia por reporte incluido y precio por reporte sobrefacturado), con activación o desactivación de cada plan.", "Media"),
        ("RF-25", "Asignación de plan a empresa", "Permite al propietario asignar o cambiar el plan de una empresa, conservando una copia de los valores contratados y registrando el cambio en el historial de planes.", "Media"),
        ("RF-26", "Generación del período de facturación", "Genera automáticamente el período de facturación mensual por empresa, con sus montos, reportes incluidos, consumidos y sobrefacturados, y su estado (pendiente, parcial, pagado, vencido o bloqueado).", "Media"),
        ("RF-27", "Control de consumo y evaluaciones excedentes", "Contabiliza cada evaluación consumida en el período, distinguiendo entre reportes incluidos en el plan y reportes sobrefacturados (excedentes), y actualiza los montos correspondientes.", "Media"),
        ("RF-28", "Registro de pagos", "Permite iniciar y registrar los pagos de un período de facturación mediante la pasarela Pagopar y un modo de prueba, generando la transacción con su referencia y estado.", "Media"),
        ("RF-29", "Confirmación de pagos", "Procesa la confirmación del pago (vía webhook de Pagopar o aprobación en modo de prueba), actualiza el saldo del período y lo marca como parcial o pagado según corresponda.", "Media"),
        ("RF-30", "Bloqueo y desbloqueo por facturación", "Bloquea las consultas de una empresa ante deuda pendiente y las desbloquea automáticamente al confirmarse el pago del período adeudado.", "Media"),
    ]),
]


def main():
    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    d.add_heading("Requerimientos Funcionales", level=1)

    nota = d.add_paragraph()
    r = nota.add_run("Nota de revisión: se reorganizó y amplió el bloque del modelo comercial "
                     "(planes, facturación y pagos) para documentarlo formalmente como parte del "
                     "sistema, no solo en la base de datos. Además, se corrigió el requerimiento de "
                     "auditoría eliminando la mención a la dirección IP, dato que el sistema no registra.")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = GRIS

    for titulo, filas in bloques:
        d.add_heading(titulo, level=2)
        tabla = d.add_table(rows=1, cols=4)
        tabla.style = "Light Grid Accent 1"
        hdr = tabla.rows[0].cells
        for c, txt in zip(hdr, ["ID", "Nombre", "Descripción", "Prioridad"]):
            c.text = ""
            run = c.paragraphs[0].add_run(txt)
            run.bold = True
        for (rid, nombre, desc, prio) in filas:
            celdas = tabla.add_row().cells
            celdas[0].text = rid
            celdas[1].text = nombre
            celdas[2].text = desc
            celdas[3].text = prio
        d.add_paragraph()

    d.save(RUTA)
    print("OK ->", RUTA)


if __name__ == "__main__":
    main()
