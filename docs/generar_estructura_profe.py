# -*- coding: utf-8 -*-
"""
Genera un Word con la ESTRUCTURA RECOMENDADA que indico el profesor,
tal cual fue entregada (sin desarrollar el contenido).
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

AZUL = RGBColor(0x10, 0x2A, 0x47)
AZUL_CLARO = RGBColor(0x2E, 0x5A, 0x8C)
GRIS = RGBColor(0x55, 0x55, 0x55)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

# ---- Título ----
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("ESTRUCTURA RECOMENDADA PARA EL APARTADO DEL LIBRO")
r.bold = True
r.font.size = Pt(18)
r.font.color.rgb = AZUL

doc.add_paragraph()

# ---- Indicación de ubicación ----
p = doc.add_paragraph()
r = p.add_run(
    "Una vez realizadas estas correcciones, incorporar dentro de \u201cModelado del Sistema\u201d, "
    "después de Arquitectura de Solución y antes del Marco Analítico, el apartado:"
)
r.font.size = Pt(11)

# ---- Nombre del apartado ----
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("\u201cModelo de Negocio y Lógica Operativa de la Solución\u201d")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = AZUL_CLARO

p = doc.add_paragraph()
r = p.add_run("Con la siguiente estructura:")
r.font.size = Pt(11)

# ---- Lista de los 12 puntos ----
puntos = [
    "Descripción del modelo comercial.",
    "Actores del modelo: cliente, usuario, pagador y beneficiario.",
    "Problema y propuesta de valor.",
    "Lógica operativa del modelo.",
    "Reglas principales del negocio.",
    "Flujo de valor.",
    "Cómo el sistema soporta el modelo de negocio.",
    "Estructura de costos.",
    "Modelo de ingresos y sostenibilidad.",
    "Business Model Canvas final.",
    "Análisis integral del Canvas.",
    "Viabilidad y coherencia del modelo.",
]

for punto in puntos:
    doc.add_paragraph(punto, style="List Number")

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Estructura_Apartado_Modelo_Negocio.docx")
doc.save(out)
print("GENERADO:", out)
