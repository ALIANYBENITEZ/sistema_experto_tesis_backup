# -*- coding: utf-8 -*-
"""
Genera el Word del Plan de Pruebas: 6 casos (uno por modulo) con flujo correcto.
Cada caso es una tabla de 2 columnas (Elemento / Descripcion).
Basado en los modulos reales del sistema (auth 2FA, usuarios, clientes,
evaluacion por criterios, scoring IF-THEN, reportes).
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x10, 0x2A, 0x47)
GRIS = RGBColor(0x55, 0x55, 0x55)

doc = Document()

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def h(text, level=1, color=AZUL):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = color
        run.font.name = "Times New Roman"
    return p


def para(text, size=11, italic=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    if color:
        r.font.color.rgb = color
    return p


def caso(titulo_modulo, filas):
    """filas: lista de (Elemento, Descripcion)."""
    cap = doc.add_paragraph()
    r = cap.add_run(titulo_modulo)
    r.bold = True
    r.font.size = Pt(11.5)
    r.font.name = "Times New Roman"
    r.font.color.rgb = AZUL
    cap.paragraph_format.space_before = Pt(10)
    cap.paragraph_format.space_after = Pt(3)

    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for elemento, descripcion in filas:
        cells = t.add_row().cells
        # Columna 1 (Elemento) con fondo
        cells[0].text = ""
        p0 = cells[0].paragraphs[0]
        r0 = p0.add_run(elemento)
        r0.bold = True
        r0.font.size = Pt(10)
        r0.font.name = "Times New Roman"
        set_cell_bg(cells[0], "E8EDF3")
        cells[0].width = Inches(1.7)
        # Columna 2 (Descripcion)
        cells[1].text = ""
        p1 = cells[1].paragraphs[0]
        r1 = p1.add_run(descripcion)
        r1.font.size = Pt(10)
        r1.font.name = "Times New Roman"
        cells[1].width = Inches(4.8)
    return t


# ---- Titulo ----
tt = doc.add_paragraph()
tt.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tt.add_run("Plan de Pruebas de los Flujos Principales")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = AZUL
r.font.name = "Times New Roman"

para(
    "El plan de pruebas verifica el funcionamiento correcto (flujo exitoso) de los módulos "
    "principales del sistema. Se define un caso de prueba por módulo, describiendo sus criterios de "
    "aceptación, el control de errores, el resultado esperado y el obtenido. La ejecución de las "
    "pruebas se apoya en PyTest sobre la API REST.",
    italic=True, color=GRIS,
)

# ---- Casos ----
caso("Caso de prueba CP-01 — Módulo de Autenticación", [
    ("Identificador", "CP-01"),
    ("Nombre", "Inicio de sesión con credenciales válidas y verificación 2FA"),
    ("Criterios de Aceptación",
     "El sistema debe autenticar al usuario que ingresa credenciales correctas y, tras validar el "
     "código OTP de Google Authenticator, otorgar acceso completo mediante tokens JWT (access y refresh)."),
    ("Control de Errores",
     "La capa de autenticación (Flask-JWT-Extended) valida las credenciales y el estado del usuario; "
     "ante un OTP incorrecto o un usuario inactivo retorna una respuesta JSON controlada (success: false) "
     "con el código HTTP correspondiente, sin exponer datos sensibles."),
    ("Resultado Esperado",
     "Acceso concedido con estado \u201cauthenticated\u201d, emisión del access token (8 h) y refresh "
     "token (30 d), y registro del evento LOGIN_EXITOSO en la auditoría."),
    ("Resultado Obtenido",
     "El sistema autenticó al usuario, validó el OTP, emitió los tokens JWT y registró el evento en la "
     "bitácora de auditoría. El flujo completo fue validado mediante PyTest."),
    ("Evidencias",
     "Ver Figura XX - Respuesta de la API de login con tokens JWT y Figura XX - Consola de PyTest con "
     "el test de autenticación superado."),
])

caso("Caso de prueba CP-02 — Módulo de Gestión de Usuarios", [
    ("Identificador", "CP-02"),
    ("Nombre", "Alta de usuario con rol válido por un administrador"),
    ("Criterios de Aceptación",
     "Un usuario con rol administrador debe poder registrar un nuevo usuario indicando nombre, "
     "apellido, email, contraseña y rol (administrador u operador), con la contraseña almacenada de "
     "forma cifrada."),
    ("Control de Errores",
     "El endpoint valida los campos obligatorios, el rol permitido y la unicidad del email; ante un "
     "email ya registrado retorna HTTP 409 con una respuesta JSON controlada (success: false) sin "
     "crear duplicados."),
    ("Resultado Esperado",
     "Creación del usuario con contraseña hasheada (scrypt), respuesta HTTP 201 con los datos del "
     "usuario y mensaje \u201cUsuario creado\u201d."),
    ("Resultado Obtenido",
     "El sistema registró el usuario correctamente, almacenó la contraseña cifrada y rechazó los "
     "emails duplicados con el mensaje controlado esperado. Validado mediante PyTest."),
    ("Evidencias",
     "Ver Figura XX - Respuesta de la API al crear un usuario y Figura XX - Consola de PyTest con el "
     "test de alta de usuario superado."),
])

caso("Caso de prueba CP-03 — Módulo de Gestión de Clientes", [
    ("Identificador", "CP-03"),
    ("Nombre", "Registro de cliente con datos válidos"),
    ("Criterios de Aceptación",
     "El sistema debe permitir registrar un cliente con tipo y número de documento válidos (CI, RUC o "
     "PAS), teléfono sin letras y fecha de nacimiento correspondiente a un mayor de edad (\u2265 18 años)."),
    ("Control de Errores",
     "La capa de validación verifica el formato del documento, el teléfono y la edad mínima; ante "
     "datos inválidos retorna una respuesta JSON controlada (success: false, HTTP 400) con el mensaje "
     "específico del error, sin persistir el registro."),
    ("Resultado Esperado",
     "Alta del cliente con estado \u201cactivo\u201d, asociación a la empresa del usuario y registro "
     "del evento en la auditoría."),
    ("Resultado Obtenido",
     "El sistema registró el cliente con datos válidos, aplicó las validaciones de documento, teléfono "
     "y edad, y rechazó los casos inválidos con los mensajes controlados esperados. Validado mediante "
     "PyTest."),
    ("Evidencias",
     "Ver Figura XX - Respuesta de la API al registrar un cliente y Figura XX - Consola de PyTest con "
     "el test de registro de cliente superado."),
])

caso("Caso de prueba CP-04 — Módulo de Evaluación por Criterios", [
    ("Identificador", "CP-04"),
    ("Nombre", "Cálculo del puntaje ponderado de una evaluación por criterios"),
    ("Criterios de Aceptación",
     "El sistema debe calcular el puntaje total (escala 0\u2013100) como la media ponderada de los "
     "criterios activos (Capacidad de Pago, Historial Crediticio, Estabilidad Laboral, Documentación y "
     "Patrimonio) según su peso."),
    ("Control de Errores",
     "El motor recorta cada valor al rango 0\u2013100 e ignora criterios inexistentes; si no hay "
     "criterios activos o el peso total es cero, retorna un resultado controlado (\u201crechazado\u201d) "
     "sin provocar error de división."),
    ("Resultado Esperado",
     "Puntaje ponderado correcto y resultado automático: aprobado (\u2265 70), observado (\u2265 50) o "
     "rechazado (< 50)."),
    ("Resultado Obtenido",
     "El sistema calculó el puntaje ponderado conforme a los pesos configurados y asignó el resultado "
     "correcto según los umbrales. El cálculo fue validado mediante PyTest con distintos perfiles."),
    ("Evidencias",
     "Ver Figura XX - Respuesta de la API con el puntaje ponderado y Figura XX - Consola de PyTest con "
     "el test de evaluación por criterios superado."),
])

caso("Caso de prueba CP-05 — Módulo de Scoring (Motor de Reglas IF-THEN)", [
    ("Identificador", "CP-05"),
    ("Nombre", "Ejecución del motor de reglas y clasificación de riesgo"),
    ("Criterios de Aceptación",
     "El sistema debe evaluar los datos del cliente contra las reglas IF-THEN del motor, sumar los "
     "puntos de las reglas cumplidas y clasificar el riesgo en bajo (\u2265 75 %), medio (\u2265 50 %) "
     "o alto (< 50 %) sobre el puntaje máximo."),
    ("Control de Errores",
     "Si falla una regla determinante (lista negra ONU/OFAC o edad mínima) el motor corta la "
     "evaluación y clasifica como \u201crechazado\u201d; los parámetros sin dato se tratan como regla "
     "no cumplida sin interrumpir el proceso."),
    ("Resultado Esperado",
     "Score final, porcentaje y clasificación de riesgo correctos, con el detalle de cada regla "
     "(cumplida/no cumplida) y una recomendación de apoyo a la decisión."),
    ("Resultado Obtenido",
     "El sistema ejecutó el motor, calculó el score y la clasificación, aplicó las reglas determinantes "
     "y registró el detalle por regla para trazabilidad. Validado mediante PyTest."),
    ("Evidencias",
     "Ver Figura XX - Respuesta de la API con el resultado del scoring y Figura XX - Consola de PyTest "
     "con el test del motor de reglas superado."),
])

caso("Caso de prueba CP-06 — Módulo de Reportes", [
    ("Identificador", "CP-06"),
    ("Nombre", "Generación de reporte de evaluaciones con filtros"),
    ("Criterios de Aceptación",
     "El sistema debe generar reportes de las evaluaciones de riesgo permitiendo filtrar por rango de "
     "fechas, clasificación de riesgo y evaluador, respetando el alcance por empresa del usuario."),
    ("Control de Errores",
     "Los parámetros de fecha inválidos se descartan de forma segura (se ignora el filtro) sin provocar "
     "error; la consulta aplica automáticamente el filtro por empresa para no exponer datos de otras "
     "organizaciones."),
    ("Resultado Esperado",
     "Listado de evaluaciones acorde a los filtros aplicados y exportación del reporte (PDF) con los "
     "datos del cliente, puntaje, clasificación y responsable."),
    ("Resultado Obtenido",
     "El sistema generó el reporte con los filtros aplicados, respetó el aislamiento por empresa y "
     "produjo la salida esperada. Validado mediante PyTest."),
    ("Evidencias",
     "Ver Figura XX - Reporte generado por la API y Figura XX - Consola de PyTest con el test de "
     "reportes superado."),
])

# ---- Nota final ----
doc.add_paragraph()
nota = doc.add_paragraph()
r = nota.add_run("Nota. ")
r.italic = True
r.font.size = Pt(10)
r.font.name = "Times New Roman"
r2 = nota.add_run(
    "Los casos describen el flujo correcto (camino exitoso) de cada módulo principal del sistema. "
    "Las referencias \u201cFigura XX\u201d deben reemplazarse por la numeración definitiva del documento. "
    "Elaboración propia (2026)."
)
r2.font.size = Pt(10)
r2.font.name = "Times New Roman"

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Plan_de_Pruebas_Flujos_Principales.docx")
doc.save(out)
print("GENERADO:", out)
