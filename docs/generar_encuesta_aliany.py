# -*- coding: utf-8 -*-
"""
Genera el Word de la ENCUESTA (P1-P20) con formato de tabla academica
estilo APA (3 columnas: codigo / pregunta / opciones), lineas horizontales,
texto centrado. Formato segun la imagen de referencia. Tesis 2026 - ALIANY.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NEGRO = RGBColor(0x00, 0x00, 0x00)
GRIS = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# Margenes
for section in doc.sections:
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.5


def _set_border(cell, edges):
    """Aplica bordes solo en los lados indicados (top/bottom) en negro."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        if edge in edges:
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "6")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "000000")
        else:
            element.set(qn("w:val"), "nil")
        tcBorders.append(element)
    tcPr.append(tcBorders)


def cell_text(cell, text, bold=False, align="center", size=12, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(4)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"


# ---- Titulo ----
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Instrumento de recolección de datos")
r.bold = True
r.font.size = Pt(13)
r.font.name = "Times New Roman"

st = doc.add_paragraph()
st.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = st.add_run("Encuesta dirigida a inmobiliarias sobre el proceso de evaluación de riesgo de clientes")
r.font.size = Pt(11)
r.italic = True
r.font.name = "Times New Roman"

# Caption de tabla (estilo APA)
cap = doc.add_paragraph()
r = cap.add_run("Tabla 1")
r.bold = True
r.font.size = Pt(11)
r.font.name = "Times New Roman"
cap2 = doc.add_paragraph()
r = cap2.add_run("Preguntas del cuestionario y opciones de respuesta")
r.italic = True
r.font.size = Pt(11)
r.font.name = "Times New Roman"
cap2.paragraph_format.space_after = Pt(2)

# ---- Datos de la encuesta ----
preguntas = [
    ("P1", "¿Su inmobiliaria cuenta con un proceso definido para evaluar el riesgo de los clientes antes de una operación de compra o financiamiento de un inmueble?",
     "Sí / No / Parcialmente"),
    ("P2", "¿Cómo realiza actualmente su empresa la evaluación del perfil de los clientes?",
     "Revisión manual de documentos / Uso de herramientas informáticas / Consulta de información externa / Otro: ______"),
    ("P3", "¿Qué nivel de automatización tiene actualmente su proceso de evaluación?",
     "Totalmente manual (papeles) / Manual con Excel / Mixto (herramientas + revisión manual) / Mayormente automatizado (software)"),
    ("P4", "¿Cuánto tiempo demora aproximadamente el análisis de un cliente antes de tomar una decisión?",
     "Menos de 10 min / 10 a 20 min / 20 a 40 min / Más de 40 min"),
    ("P5", "¿Cuáles son las principales limitaciones del proceso actual? (marque todas las que correspondan)",
     "Errores humanos / Información desorganizada / Reprocesos por datos incompletos / Pérdida de documentos / Falta de criterio uniforme entre evaluadores / Demora en información externa / Otro: ______"),
    ("P7", "¿Su empresa utiliza criterios definidos o reglas establecidas para clasificar el nivel de riesgo?",
     "Sí / No / Parcialmente"),
    ("P8", "¿Con qué frecuencia tiene dificultades para analizar la información por la cantidad de datos o falta de organización?",
     "Siempre / Frecuentemente / A veces / Rara vez / Nunca"),
    ("P9", "La evaluación manual puede generar resultados diferentes según la persona encargada del análisis.",
     "Totalmente de acuerdo / De acuerdo / Ni de acuerdo ni en desacuerdo / En desacuerdo / Totalmente en desacuerdo"),
    ("P10", "¿Qué nivel de importancia tiene para usted contar con una clasificación automática del nivel de riesgo del cliente?",
     "Muy importante / Importante / Moderadamente importante / Poco importante / Nada importante"),
    ("P11", "¿Considera necesario un sistema que registre y centralice la información de los clientes evaluados?",
     "Sí / No / Parcialmente"),
    ("P12", "Un sistema experto que genere automáticamente un puntaje de scoring permitiría mejorar el proceso actual.",
     "Totalmente de acuerdo / De acuerdo / Ni de acuerdo ni en desacuerdo / En desacuerdo / Totalmente en desacuerdo"),
    ("P13", "La clasificación automática en niveles de riesgo (bajo, medio y alto) facilitaría la toma de decisiones comerciales.",
     "Totalmente de acuerdo / De acuerdo / Ni de acuerdo ni en desacuerdo / En desacuerdo / Totalmente en desacuerdo"),
    ("P14", "¿Qué tipos de reportes le resultarían más útiles? (marque todas las que correspondan)",
     "Ficha individual de riesgo por cliente / Historial de evaluaciones / Comparativo o ranking por nivel de riesgo / Estadísticas generales / Reporte por período / Exportable en PDF / Otro: ______"),
    ("P15", "¿Qué datos considera imprescindibles en el reporte de un cliente? (marque todas las que correspondan)",
     "Puntaje total de scoring / Nivel de riesgo (bajo/medio/alto) / Detalle por criterio / Justificación de la decisión / Recomendación final / Fecha y responsable"),
    ("P16", "¿Cuánto tiempo considera aceptable que tome el sistema en entregar un resultado?",
     "Menos de 1 min / 1 a 3 min / 3 a 5 min / Más de 5 min"),
    ("P17", "¿Qué tan importante es que el sistema entregue resultados consistentes (mismo cliente = mismo puntaje, sin importar quién lo evalúe)?",
     "Muy importante / Importante / Moderadamente importante / Poco importante / Nada importante"),
    ("P18", "¿Es importante que el sistema controle el acceso de usuarios para proteger la información de los clientes?",
     "Muy importante / Importante / Moderadamente importante / Poco importante / Nada importante"),
    ("P19", "¿Qué tan importante es que el sistema permita consultar antecedentes o información complementaria para apoyar la evaluación?",
     "Muy importante / Importante / Moderadamente importante / Poco importante / Nada importante"),
    ("P20", "¿Estaría dispuesto a utilizar un sistema experto web de scoring para apoyar la evaluación del riesgo de clientes en su inmobiliaria?",
     "Definitivamente sí / Probablemente sí / Indeciso / Probablemente no / Definitivamente no"),
]

# ---- Tabla principal (3 columnas) ----
t = doc.add_table(rows=0, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
col_w = [Inches(0.7), Inches(2.9), Inches(2.9)]

# Encabezado
hdr = t.add_row().cells
cell_text(hdr[0], "Código", bold=True)
cell_text(hdr[1], "Pregunta", bold=True)
cell_text(hdr[2], "Opciones de respuesta", bold=True)
for i, c in enumerate(hdr):
    c.width = col_w[i]
    _set_border(c, ("top", "bottom"))

# Insertar filas; cuando toca P6 (matriz) se intercala despues de P5
def add_fila(codigo, pregunta, opciones):
    cells = t.add_row().cells
    cell_text(cells[0], codigo)
    cell_text(cells[1], pregunta)
    cell_text(cells[2], opciones)
    for i, c in enumerate(cells):
        c.width = col_w[i]
        _set_border(c, ("bottom",))

for codigo, pregunta, opciones in preguntas:
    add_fila(codigo, pregunta, opciones)
    if codigo == "P5":
        # P6 es una pregunta de matriz: ocupa su propia descripcion y
        # se detalla en una tabla aparte debajo.
        add_fila(
            "P6",
            "Califique la importancia de cada criterio para evaluar el riesgo "
            "(escala 1 = nada importante a 5 = muy importante). Ver matriz de la Tabla 2.",
            "Escala Likert 1 a 5 por criterio",
        )

# ---- Tabla 2: matriz de criterios (P6) ----
doc.add_paragraph()
cap = doc.add_paragraph()
r = cap.add_run("Tabla 2")
r.bold = True
r.font.size = Pt(11)
r.font.name = "Times New Roman"
cap2 = doc.add_paragraph()
r = cap2.add_run("P6. Importancia de cada criterio de evaluación (escala 1 = nada importante, 5 = muy importante)")
r.italic = True
r.font.size = Pt(11)
r.font.name = "Times New Roman"
cap2.paragraph_format.space_after = Pt(2)

criterios = [
    "Ingresos y capacidad de pago",
    "Estabilidad laboral",
    "Historial de pagos",
    "Nivel de endeudamiento",
    "Referencias personales/comerciales",
    "Historial crediticio",
]

mt = doc.add_table(rows=0, cols=6)
mt.alignment = WD_TABLE_ALIGNMENT.CENTER
mcol_w = [Inches(3.0), Inches(0.6), Inches(0.6), Inches(0.6), Inches(0.6), Inches(0.6)]

mhdr = mt.add_row().cells
cell_text(mhdr[0], "Criterio", bold=True, align="left")
for i, n in enumerate(["1", "2", "3", "4", "5"], start=1):
    cell_text(mhdr[i], n, bold=True)
for i, c in enumerate(mhdr):
    c.width = mcol_w[i]
    _set_border(c, ("top", "bottom"))

for crit in criterios:
    cells = mt.add_row().cells
    cell_text(cells[0], crit, align="left")
    for i in range(1, 6):
        cell_text(cells[i], "☐")
    for i, c in enumerate(cells):
        c.width = mcol_w[i]
        _set_border(c, ("bottom",))

# ---- Nota al pie ----
doc.add_paragraph()
nota = doc.add_paragraph()
r = nota.add_run("Nota. ")
r.italic = True
r.font.size = Pt(10)
r.font.name = "Times New Roman"
r2 = nota.add_run(
    "Instrumento elaborado para la recolección de datos de la investigación. "
    "Las preguntas P5, P14 y P15 admiten selección múltiple; P6 utiliza una escala "
    "Likert de 1 a 5 por criterio. Fuente: elaboración propia (tesis 2026)."
)
r2.font.size = Pt(10)
r2.font.name = "Times New Roman"

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Encuesta_Scoring_Inmobiliarias_ALIANY.docx")
doc.save(out)
print("GENERADO:", out)
