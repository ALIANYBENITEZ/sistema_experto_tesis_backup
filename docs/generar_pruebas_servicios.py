# -*- coding: utf-8 -*-
"""
Genera el Word de PRUEBAS DE SERVICIOS (API REST) que acompaña la corrida de
Postman (colección de evidencia). Tabla de casos con método, endpoint,
resultado esperado y validación.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x10, 0x2A, 0x47)
GRIS = RGBColor(0x55, 0x55, 0x55)

doc = Document()
for s in doc.sections:
    s.left_margin = Inches(0.8)
    s.right_margin = Inches(0.8)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def h(text, level=1, color=AZUL):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = color
        run.font.name = "Times New Roman"
    return p


def para(text, size=11, italic=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(size)
    r.font.name = "Times New Roman"
    if color:
        r.font.color.rgb = color
    return p


def tabla(headers, rows, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, head in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(head)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9.5)
        run.font.name = "Times New Roman"
        set_cell_bg(hdr[i], "102A47")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            run = p.add_run(str(val))
            run.font.size = Pt(9)
            run.font.name = "Times New Roman"
            if ri % 2 == 1:
                set_cell_bg(cells[ci], "EEF2F7")
    if col_widths:
        for row in t.rows:
            for ci, w in enumerate(col_widths):
                row.cells[ci].width = Inches(w)
    return t


# ---- Título ----
tt = doc.add_paragraph()
tt.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tt.add_run("Pruebas de Servicios de la API REST")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = AZUL
r.font.name = "Times New Roman"

para(
    "Esta sección documenta las pruebas realizadas sobre los servicios web (API REST) del sistema, "
    "ejecutadas con Postman contra el entorno desplegado en OnRender "
    "(https://inmobiliaria-backend-u6z7.onrender.com). Las pruebas abarcan la validación de seguridad "
    "del inicio de sesión, la autenticación válida con doble factor (2FA) y los endpoints de los "
    "flujos principales que recuperan información y ejecutan el scoring. No se incluyen operaciones de "
    "modificación o eliminación (PUT, DELETE, PATCH) en esta corrida de evidencia.",
    italic=True, color=GRIS,
)

# ── Validación de seguridad del login ──
h("1. Validación de seguridad del inicio de sesión", 2)
para("Pruebas negativas que verifican que el sistema rechaza accesos no autorizados y solicitudes mal "
     "formadas, protegiendo el acceso a la información.")
tabla(
    ["ID", "Caso de prueba", "Método / Endpoint", "Resultado esperado"],
    [
        ["SEG-01", "Login con contraseña incorrecta",
         "POST /api/auth/login", "HTTP 401. No se entrega token. success distinto de true."],
        ["SEG-02", "Login sin credenciales (cuerpo vacío)",
         "POST /api/auth/login", "HTTP 400. Mensaje de datos inválidos."],
        ["SEG-03", "Acceso a endpoint protegido sin token",
         "GET /api/auth/me", "HTTP 401. Acceso bloqueado."],
        ["SEG-04", "Acceso con token manipulado o inválido",
         "GET /api/clients/", "HTTP 401 o 422. Token rechazado."],
    ],
    col_widths=[0.8, 2.3, 2.0, 2.4],
)
para("Nota: El sistema utiliza autenticación JWT (Flask-JWT-Extended) con doble factor obligatorio "
     "(TOTP / Google Authenticator). Las contraseñas se almacenan cifradas con scrypt.",
     size=9.5, italic=True, color=GRIS)

# ── Autenticación válida ──
h("2. Autenticación válida (login + 2FA)", 2)
para("Flujo exitoso de autenticación de un usuario válido, incluyendo la verificación del segundo "
     "factor y la emisión de los tokens de acceso.")
tabla(
    ["ID", "Caso de prueba", "Método / Endpoint", "Resultado esperado"],
    [
        ["AUT-01", "Inicio de sesión con credenciales correctas",
         "POST /api/auth/login", "HTTP 200. Devuelve auth_state y partial_token."],
        ["AUT-02", "Verificación del segundo factor (OTP)",
         "POST /api/auth/2fa/verify", "HTTP 200. auth_state = authenticated. Emite access y refresh token."],
        ["AUT-03", "Consulta del usuario autenticado",
         "GET /api/auth/me", "HTTP 200. Devuelve email y rol del usuario."],
    ],
    col_widths=[0.8, 2.3, 2.0, 2.4],
)

# ── Flujos principales (lectura) ──
h("3. Flujos principales de recuperación de información", 2)
para("Endpoints de consulta que sostienen los flujos principales del sistema. Verifican la respuesta "
     "normalizada de la API y el aislamiento de datos por empresa.")
tabla(
    ["ID", "Caso de prueba", "Método / Endpoint", "Resultado esperado"],
    [
        ["FLU-01", "Verificar estado del servicio",
         "GET /api/health", "HTTP 200. status = ok."],
        ["FLU-02", "Listar clientes (paginado)",
         "GET /api/clients/", "HTTP 200. Respuesta con items (arreglo)."],
        ["FLU-03", "Listar criterios de evaluación",
         "GET /api/evaluations/criteria", "HTTP 200. Arreglo de criterios con peso."],
        ["FLU-04", "Obtener parámetros del scoring",
         "GET /api/scoring/parametros", "HTTP 200. success = true."],
        ["FLU-05", "Listar motores de reglas",
         "GET /api/scoring/motores", "HTTP 200. Arreglo de motores."],
        ["FLU-06", "Listar evaluaciones de riesgo",
         "GET /api/scoring/evaluaciones", "HTTP 200. success = true."],
        ["FLU-07", "Resumen del dashboard de reportes",
         "GET /api/reports/summary", "HTTP 200. Objeto con estadísticas."],
        ["FLU-08", "Consultar bitácora de auditoría",
         "GET /api/auditoria/", "HTTP 200. Registros de auditoría."],
    ],
    col_widths=[0.8, 2.3, 2.1, 2.3],
)

# ── Flujo principal: scoring ──
h("4. Flujo principal: ejecución del scoring", 2)
para("Prueba del flujo central del sistema: la evaluación de riesgo de un cliente mediante el motor "
     "de reglas IF-THEN.")
tabla(
    ["ID", "Caso de prueba", "Método / Endpoint", "Resultado esperado"],
    [
        ["SCO-01", "Evaluar un cliente (ejecutar el motor IF-THEN)",
         "POST /api/scoring/evaluar",
         "HTTP 200. Devuelve categoria_riesgo (bajo / medio / alto / rechazado) y el detalle por regla."],
    ],
    col_widths=[0.8, 2.3, 2.0, 2.4],
)

# ── Procedimiento de ejecución ──
h("5. Procedimiento de ejecución de las pruebas", 2)
doc.add_paragraph("Importar en Postman la colección \u201cInmobiliaria - Pruebas de Servicios (Evidencia)\u201d "
                  "y el entorno \u201cInmobiliaria - OnRender\u201d.", style="List Number")
doc.add_paragraph("Cargar el secreto TOTP del usuario en la variable totp_secret del entorno (una sola vez).",
                  style="List Number")
doc.add_paragraph("Ejecutar la colección con el Collection Runner, respetando el orden de las carpetas "
                  "(Seguridad, Autenticación, Flujos de lectura, Scoring).", style="List Number")
doc.add_paragraph("Verificar en el reporte del Runner que todos los casos resulten aprobados (passed).",
                  style="List Number")

# ---- Nota final ----
doc.add_paragraph()
nota = doc.add_paragraph()
r = nota.add_run("Nota. ")
r.italic = True
r.font.size = Pt(10)
r.font.name = "Times New Roman"
r2 = nota.add_run(
    "Las pruebas se ejecutaron sobre el servicio desplegado en OnRender mediante Postman, con "
    "verificación automática del segundo factor (TOTP) y validación del cuerpo de cada respuesta. "
    "El entorno gratuito de OnRender puede presentar una latencia inicial mayor al reactivarse el "
    "servicio. Elaboración propia (2026)."
)
r2.font.size = Pt(10)
r2.font.name = "Times New Roman"

import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Pruebas_de_Servicios_API.docx")
doc.save(out)
print("GENERADO:", out)
