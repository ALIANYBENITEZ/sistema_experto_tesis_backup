# -*- coding: utf-8 -*-
"""Diccionario de Datos en WORD (.docx) con descripcion enriquecida por campo."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from app import create_app
from app.extensions import db
from sqlalchemy import text
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"c:\Users\RYZEN\Documents\TESIS 2026\SISTEMA\V.0.1\DICCIONARIO_DE_DATOS.docx"

# Reusar descripciones del PDF (import de un modulo temporal seria ideal;
# aqui las definimos de nuevo, resumidas, cargando desde el mismo dict).
DESC_TABLA = {
    "usuarios":"Almacena la informacion principal de las personas con acceso al sistema (propietario, administradores y comerciales) junto con sus parametros de seguridad y autenticacion de dos factores.",
    "empresas":"Registra las empresas (inmobiliarias) que contratan el uso del sistema, incluyendo su estado y el control de bloqueo de consultas por facturacion.",
    "clientes":"Almacena los datos de los clientes (personas fisicas o juridicas) que son evaluados por el sistema.",
    "cliente_empresa":"Tabla de relacion que vincula clientes con empresas (muchos a muchos). Permite que un mismo cliente este registrado en varias empresas.",
    "paises":"Catalogo de paises utilizado para la nacionalidad de los clientes.",
    "departamento":"Catalogo de departamentos de Paraguay.",
    "ciudad":"Catalogo de ciudades, relacionadas a un departamento.",
    "documentos":"Almacena los documentos adjuntos de cada cliente.",
    "historial_crediticio":"Registra los datos financieros e historicos de cada cliente para la evaluacion de riesgo. Relacion uno a uno con clientes.",
    "criterios":"Catalogo de criterios ponderados utilizados en la evaluacion por criterios.",
    "evaluaciones":"Registra las evaluaciones por criterios ponderados realizadas a los clientes.",
    "evaluacion_detalle":"Detalle de cada criterio evaluado dentro de una evaluacion por criterios.",
    "motor_reglas":"Define los motores de reglas (sistema experto) asignados a cada empresa.",
    "reglas":"Reglas individuales de tipo IF-THEN que componen un motor de reglas.",
    "evaluacion_riesgo":"Registra el resultado de las evaluaciones de riesgo realizadas mediante el motor de reglas.",
    "resultado_detalle":"Detalle por regla de una evaluacion de riesgo.",
    "detalle_operacion":"Datos de la operacion inmobiliaria asociada a una evaluacion de riesgo. Relacion uno a uno.",
    "scoring_modelo":"Define los modelos de scoring configurables del Core, asignados a cada empresa.",
    "scoring_factor":"Factores de evaluacion que componen un modelo de scoring.",
    "scoring_catalogo":"Valores posibles (catalogo) para factores de tipo categorico.",
    "scoring_regla":"Reglas IF-THEN asociadas a cada factor del modelo de scoring.",
    "scoring_umbral":"Umbrales de clasificacion de riesgo por modelo de scoring.",
    "scoring_evaluacion":"Registra el resultado de las evaluaciones realizadas con el Core de scoring.",
    "scoring_detalle":"Detalle del resultado por cada factor evaluado en una evaluacion del Core.",
    "planes":"Catalogo de planes comerciales prepago (tabla parametrica de facturacion).",
    "empresa_planes":"Plan vigente contratado por cada empresa, con copia de valores al contratar.",
    "periodos_facturacion":"Periodos mensuales de facturacion por empresa, con detalle de consumo y montos.",
    "consumo_reportes":"Registro individual de cada reporte consumido por una empresa.",
    "pagos":"Registra las transacciones de pago de las empresas.",
    "historial_empresa_planes":"Historial de cambios de plan por empresa.",
    "auditoria":"Bitacora centralizada que registra los eventos y acciones realizadas en el sistema.",
    "lista_negra_onu":"Tabla de referencia con la lista consolidada de sanciones ONU/OFAC.",
}

DESC_CAMPO = {
 "usuarios.id":"Identificador unico del usuario.","usuarios.nombre":"Nombre de pila del usuario.",
 "usuarios.apellido":"Apellido del usuario.","usuarios.email":"Correo electronico usado para iniciar sesion. Debe ser unico.",
 "usuarios.password_hash":"Contrasena almacenada de forma cifrada (hash).","usuarios.rol":"Rol del usuario: propietario, administrador o comercial.",
 "usuarios.activo":"Indica si el usuario esta habilitado (1) o inhabilitado (0).","usuarios.creado_en":"Fecha y hora de creacion del registro.",
 "usuarios.actualizado_en":"Fecha y hora de la ultima modificacion.","usuarios.id_empresa":"Empresa a la que pertenece el usuario (0 = propietario).",
 "usuarios.totp_secret":"Clave secreta para la autenticacion de dos factores (2FA).","usuarios.totp_estado":"Estado del 2FA: no_configurado, pendiente, activado o restablecido.",
 "usuarios.totp_fecha_config":"Fecha en que se configuro el 2FA.","usuarios.totp_fecha_reset":"Fecha del ultimo restablecimiento del 2FA.",
 "empresas.id":"Identificador unico de la empresa.","empresas.nombre":"Razon social o nombre de la empresa.",
 "empresas.ruc":"Registro unico del contribuyente. Debe ser unico.","empresas.direccion":"Direccion fisica de la empresa.",
 "empresas.telefono":"Telefono de contacto.","empresas.email":"Correo electronico de contacto.",
 "empresas.activo":"Indica si la empresa esta activa (1) o inactiva (0).","empresas.creado_en":"Fecha y hora de creacion del registro.",
 "empresas.consultas_habilitadas":"Indica si la empresa puede realizar evaluaciones (1) o esta bloqueada (0).",
 "empresas.motivo_bloqueo":"Motivo del bloqueo: BLOQUEADO_DEUDA o BLOQUEADO_MANUAL.","empresas.fecha_bloqueo":"Fecha en que se bloquearon las consultas.",
 "empresas.fecha_desbloqueo":"Fecha en que se desbloquearon las consultas.",
 "clientes.id":"Identificador unico del cliente.","clientes.tipo_doc":"Tipo de documento: CI, RUC o PAS.",
 "clientes.num_doc":"Numero de documento. Debe ser unico.","clientes.nombre":"Nombre del cliente.","clientes.apellido":"Apellido del cliente.",
 "clientes.email":"Correo electronico del cliente.","clientes.telefono":"Telefono de contacto.","clientes.direccion":"Direccion del cliente.",
 "clientes.fecha_nacimiento":"Fecha de nacimiento del cliente.","clientes.estado":"Estado del cliente: activo o inactivo.",
 "clientes.creado_por":"Usuario que registro al cliente.","clientes.creado_en":"Fecha y hora de creacion del registro.",
 "clientes.nacionalidad":"Pais de nacionalidad del cliente.","clientes.id_ciudad":"Ciudad de residencia del cliente.",
 "cliente_empresa.id":"Identificador unico del vinculo.","cliente_empresa.id_cliente":"Cliente vinculado.",
 "cliente_empresa.id_empresa":"Empresa vinculada.","cliente_empresa.estado":"Estado del vinculo: activo o inactivo.","cliente_empresa.creado_en":"Fecha y hora de creacion del vinculo.",
 "paises.id_pais":"Codigo del pais (ej: PY, AR). Clave primaria.","paises.nombre_pais":"Nombre del pais.",
 "departamento.id_departamento":"Identificador unico del departamento.","departamento.nombre_departamento":"Nombre del departamento.",
 "ciudad.id_ciudad":"Identificador unico de la ciudad.","ciudad.nombre_ciudad":"Nombre de la ciudad.","ciudad.id_departamento":"Departamento al que pertenece la ciudad.",
 "documentos.id":"Identificador unico del documento.","documentos.cliente_id":"Cliente propietario del documento.","documentos.tipo":"Tipo de documento (ej: cedula, contrato).",
 "documentos.nombre_archivo":"Nombre original del archivo.","documentos.ruta":"Ruta de almacenamiento del archivo.","documentos.subido_por":"Usuario que subio el documento.","documentos.subido_en":"Fecha y hora de la carga.",
 "historial_crediticio.id":"Identificador unico del registro.","historial_crediticio.cliente_id":"Cliente asociado. Unico (relacion 1 a 1).",
 "historial_crediticio.fuente_externa":"Fuente de la informacion crediticia.","historial_crediticio.score_externo":"Puntaje crediticio de una fuente externa.",
 "historial_crediticio.cantidad_atrasos":"Cantidad de atrasos en pagos registrados.","historial_crediticio.deuda_total_sistema":"Deuda total del cliente en el sistema financiero.",
 "historial_crediticio.historial_pagos":"Calificacion del historial: bueno, regular, malo, sin_historial.","historial_crediticio.en_lista_negra":"Indica si el cliente esta en lista negra (1) o no (0).",
 "historial_crediticio.nivel_endeudamiento":"Porcentaje de endeudamiento del cliente.","historial_crediticio.meses_empleo_actual":"Antiguedad laboral en meses.",
 "historial_crediticio.tipo_empleo":"Tipo de empleo: dependiente, independiente, desempleado.","historial_crediticio.referencias_personales":"Calificacion de referencias personales.",
 "historial_crediticio.fecha_consulta":"Fecha de consulta de los datos.","historial_crediticio.actualizado_en":"Fecha de la ultima actualizacion.",
 "criterios.id":"Identificador unico del criterio.","criterios.nombre":"Nombre del criterio de evaluacion.","criterios.descripcion":"Descripcion del criterio.",
 "criterios.peso":"Porcentaje de ponderacion del criterio.","criterios.activo":"Indica si el criterio esta activo (1) o no (0).",
 "evaluaciones.id":"Identificador unico de la evaluacion.","evaluaciones.cliente_id":"Cliente evaluado.","evaluaciones.evaluador_id":"Usuario que realizo la evaluacion.",
 "evaluaciones.fecha":"Fecha y hora de la evaluacion.","evaluaciones.puntaje_total":"Puntaje total obtenido.","evaluaciones.resultado":"Resultado: aprobado, observado o rechazado.",
 "evaluaciones.observaciones":"Observaciones de la evaluacion.","evaluaciones.estado":"Estado de la evaluacion.",
 "evaluacion_detalle.id":"Identificador unico del detalle.","evaluacion_detalle.evaluacion_id":"Evaluacion a la que pertenece.","evaluacion_detalle.criterio_id":"Criterio evaluado.",
 "evaluacion_detalle.valor":"Valor asignado al criterio.","evaluacion_detalle.comentario":"Comentario sobre el criterio evaluado.",
 "motor_reglas.id":"Identificador unico del motor.","motor_reglas.nombre":"Nombre del motor de reglas.","motor_reglas.version":"Version del motor.",
 "motor_reglas.descripcion":"Descripcion del motor.","motor_reglas.activo":"Indica si el motor esta activo (1) o no (0).","motor_reglas.creado_por":"Usuario que creo el motor.",
 "motor_reglas.creado_en":"Fecha y hora de creacion.","motor_reglas.id_empresa_motor":"Empresa propietaria del motor (0 = global).",
 "reglas.id":"Identificador unico de la regla.","reglas.motor_id":"Motor al que pertenece la regla.","reglas.nombre":"Nombre de la regla.","reglas.descripcion":"Descripcion de la regla.",
 "reglas.parametro":"Campo del cliente a evaluar.","reglas.operador":"Operador de comparacion (>=, <=, ==, etc.).","reglas.valor_referencia":"Valor umbral de comparacion.",
 "reglas.tipo_valor":"Tipo de valor: numerico, texto o booleano.","reglas.peso_puntos":"Puntos asignados si se cumple la regla.","reglas.es_determinante":"Indica si al fallar produce rechazo directo (1).",
 "reglas.activo":"Indica si la regla esta activa (1) o no (0).","reglas.orden":"Orden de evaluacion de la regla.",
 "evaluacion_riesgo.id":"Identificador unico de la evaluacion.","evaluacion_riesgo.cliente_id":"Cliente evaluado.","evaluacion_riesgo.motor_id":"Motor de reglas utilizado.",
 "evaluacion_riesgo.usuario_id":"Usuario que ejecuto la evaluacion.","evaluacion_riesgo.fecha_analisis":"Fecha y hora del analisis.","evaluacion_riesgo.score_final":"Puntaje final obtenido.",
 "evaluacion_riesgo.score_maximo":"Puntaje maximo posible.","evaluacion_riesgo.categoria_riesgo":"Clasificacion: bajo, medio, alto o rechazado.","evaluacion_riesgo.estado":"Estado de la evaluacion.",
 "evaluacion_riesgo.observaciones":"Observaciones del analisis.","evaluacion_riesgo.regla_determinante_id":"Regla que determino el resultado.","evaluacion_riesgo.id_empresa":"Empresa que realizo la evaluacion.",
 "resultado_detalle.id":"Identificador unico del detalle.","resultado_detalle.evaluacion_id":"Evaluacion de riesgo asociada.","resultado_detalle.regla_id":"Regla evaluada.",
 "resultado_detalle.cumplido":"Indica si la regla se cumplio (1) o no (0).","resultado_detalle.valor_evaluado":"Valor real del cliente al momento de evaluar.","resultado_detalle.puntos_obtenidos":"Puntos obtenidos por esta regla.",
 "detalle_operacion.id":"Identificador unico.","detalle_operacion.evaluacion_id":"Evaluacion asociada.","detalle_operacion.tipo_propiedad":"Tipo de propiedad: casa, departamento, terreno, etc.",
 "detalle_operacion.valor_propiedad":"Valor de la propiedad.","detalle_operacion.monto_solicitado":"Monto solicitado por el cliente.","detalle_operacion.plazo_meses":"Plazo de la operacion en meses.",
 "detalle_operacion.ubicacion":"Ubicacion de la propiedad.","detalle_operacion.destino":"Destino: vivienda, inversion o comercial.",
 "scoring_modelo.id":"Identificador unico del modelo.","scoring_modelo.nombre":"Nombre del modelo de scoring.","scoring_modelo.version":"Version del modelo.",
 "scoring_modelo.descripcion":"Descripcion del modelo.","scoring_modelo.id_empresa":"Empresa propietaria del modelo.","scoring_modelo.activo":"Indica si el modelo esta activo (1) o no (0).","scoring_modelo.creado_en":"Fecha y hora de creacion.",
 "scoring_factor.id":"Identificador unico del factor.","scoring_factor.modelo_id":"Modelo al que pertenece.","scoring_factor.codigo":"Codigo interno del factor.","scoring_factor.nombre":"Nombre del factor.",
 "scoring_factor.descripcion":"Descripcion del factor.","scoring_factor.tipo_dato":"Tipo de dato esperado del factor.","scoring_factor.tipo_persona":"Aplicabilidad: PF, PJ o AMBOS.",
 "scoring_factor.categoria":"Categoria: principal o complementario.","scoring_factor.obligatorio":"Indica si el factor es obligatorio (1) o no (0).","scoring_factor.activo":"Indica si el factor esta activo (1) o no (0).","scoring_factor.orden":"Orden de presentacion del factor.",
 "scoring_catalogo.id":"Identificador unico.","scoring_catalogo.factor_id":"Factor al que pertenece.","scoring_catalogo.valor":"Valor interno de la opcion.","scoring_catalogo.etiqueta":"Etiqueta visible de la opcion.","scoring_catalogo.orden":"Orden de presentacion de la opcion.",
 "scoring_regla.id":"Identificador unico de la regla.","scoring_regla.factor_id":"Factor al que pertenece.","scoring_regla.nombre":"Nombre de la regla.","scoring_regla.operador":"Operador de comparacion.",
 "scoring_regla.valor_min":"Valor minimo del rango.","scoring_regla.valor_max":"Valor maximo del rango.","scoring_regla.nivel":"Nivel de riesgo resultante: bajo, medio o alto.","scoring_regla.peso":"Peso (positivo o negativo) aplicado al score.",
 "scoring_regla.explicacion":"Explicacion de la regla.","scoring_regla.orden":"Orden de evaluacion.",
 "scoring_umbral.id":"Identificador unico.","scoring_umbral.modelo_id":"Modelo al que pertenece.","scoring_umbral.nivel":"Nivel de clasificacion: bajo, medio o alto.","scoring_umbral.score_min":"Puntaje minimo del rango.",
 "scoring_umbral.score_max":"Puntaje maximo del rango.","scoring_umbral.descripcion":"Descripcion del umbral.",
 "scoring_evaluacion.id":"Identificador unico de la evaluacion.","scoring_evaluacion.cliente_id":"Cliente evaluado.","scoring_evaluacion.modelo_id":"Modelo utilizado.","scoring_evaluacion.modelo_version":"Version del modelo al momento de evaluar.",
 "scoring_evaluacion.usuario_id":"Usuario que ejecuto la evaluacion.","scoring_evaluacion.id_empresa":"Empresa que realizo la evaluacion.","scoring_evaluacion.tipo_persona":"Tipo de persona evaluada: PF o PJ.","scoring_evaluacion.fecha":"Fecha y hora de la evaluacion.",
 "scoring_evaluacion.score_total":"Puntaje total obtenido.","scoring_evaluacion.clasificacion":"Clasificacion final: bajo, medio o alto.","scoring_evaluacion.explicacion":"Explicacion generada por el sistema.","scoring_evaluacion.factores_evaluados":"Cantidad de factores evaluados.",
 "scoring_evaluacion.factores_sin_dato":"Cantidad de factores sin dato.","scoring_evaluacion.estado":"Estado: completa o incompleta.",
 "scoring_detalle.id":"Identificador unico del detalle.","scoring_detalle.evaluacion_id":"Evaluacion asociada.","scoring_detalle.factor_codigo":"Codigo del factor evaluado.","scoring_detalle.factor_nombre":"Nombre del factor evaluado.",
 "scoring_detalle.categoria":"Categoria del factor.","scoring_detalle.valor_original":"Valor que tenia el cliente en ese factor.","scoring_detalle.estado":"Estado: evaluado, sin_dato, no_aplica o invalido.","scoring_detalle.regla_aplicada":"Nombre de la regla que se aplico.",
 "scoring_detalle.nivel":"Nivel de riesgo resultante del factor.","scoring_detalle.peso":"Peso aplicado al score.","scoring_detalle.explicacion":"Explicacion del resultado del factor.",
 "planes.id":"Identificador unico del plan.","planes.nombre":"Nombre del plan (Bronce, Plata, Oro). Unico.","planes.cantidad_reportes_incluidos":"Cantidad de reportes incluidos en el plan.","planes.precio_plan":"Precio del plan en guaranies.",
 "planes.precio_reporte_incluido":"Precio de referencia por reporte incluido.","planes.precio_reporte_sobre_facturado":"Precio de cada reporte sobre-facturado.","planes.activo":"Indica si el plan esta activo (1) o no (0).","planes.creado_en":"Fecha y hora de creacion.","planes.actualizado_en":"Fecha de la ultima modificacion.",
 "empresa_planes.id":"Identificador unico.","empresa_planes.empresa_id":"Empresa que contrata el plan.","empresa_planes.plan_id":"Plan contratado.","empresa_planes.fecha_inicio":"Fecha de inicio del plan.","empresa_planes.fecha_fin":"Fecha de finalizacion del plan.",
 "empresa_planes.estado":"Estado: activo o finalizado.","empresa_planes.precio_plan_contratado":"Precio del plan al momento de contratar.","empresa_planes.cantidad_reportes_incluidos":"Reportes incluidos contratados.","empresa_planes.precio_reporte_sobre_facturado":"Precio del reporte extra contratado.","empresa_planes.creado_en":"Fecha y hora de creacion del registro.",
 "periodos_facturacion.id":"Identificador unico del periodo.","periodos_facturacion.empresa_id":"Empresa del periodo.","periodos_facturacion.empresa_plan_id":"Plan vigente en el periodo.","periodos_facturacion.anio":"Anio del periodo.","periodos_facturacion.mes":"Mes del periodo.",
 "periodos_facturacion.fecha_inicio":"Fecha de inicio del periodo.","periodos_facturacion.fecha_fin":"Fecha de fin del periodo.","periodos_facturacion.cantidad_reportes_incluidos":"Reportes incluidos en el periodo.","periodos_facturacion.reportes_consumidos":"Cantidad de reportes consumidos.",
 "periodos_facturacion.reportes_sobre_facturados":"Cantidad de reportes sobre-facturados.","periodos_facturacion.monto_plan":"Monto correspondiente al plan.","periodos_facturacion.monto_sobre_facturado":"Monto por reportes sobre-facturados.","periodos_facturacion.monto_total":"Monto total del periodo.",
 "periodos_facturacion.monto_pagado":"Monto ya pagado.","periodos_facturacion.saldo_pendiente":"Saldo pendiente de pago.","periodos_facturacion.estado":"Estado: PENDIENTE, PARCIAL, PAGADO, VENCIDO o BLOQUEADO.",
 "consumo_reportes.id":"Identificador unico.","consumo_reportes.empresa_id":"Empresa que consume el reporte.","consumo_reportes.usuario_id":"Usuario que genero el reporte.","consumo_reportes.evaluacion_id":"Evaluacion asociada al reporte.",
 "consumo_reportes.periodo_facturacion_id":"Periodo de facturacion.","consumo_reportes.tipo_consumo":"Tipo: INCLUIDO o SOBRE_FACTURADO.","consumo_reportes.precio_unitario":"Precio unitario del reporte.","consumo_reportes.fecha_consumo":"Fecha y hora del consumo.",
 "pagos.id":"Identificador unico del pago.","pagos.empresa_id":"Empresa que realiza el pago.","pagos.periodo_facturacion_id":"Periodo que se paga.","pagos.monto":"Monto del pago.","pagos.proveedor":"Proveedor de pago: TEST, BANCARD, PAGOPAR.",
 "pagos.referencia_externa":"Referencia externa de la transaccion.","pagos.estado":"Estado: PENDIENTE, APROBADO, RECHAZADO o CANCELADO.","pagos.metodo_pago":"Metodo de pago utilizado.","pagos.observacion":"Observaciones del pago.","pagos.fecha_inicio":"Fecha de inicio de la transaccion.","pagos.fecha_confirmacion":"Fecha de confirmacion del pago.",
 "historial_empresa_planes.id":"Identificador unico.","historial_empresa_planes.empresa_id":"Empresa asociada.","historial_empresa_planes.plan_id":"Plan anterior.","historial_empresa_planes.plan_nombre":"Nombre del plan anterior.","historial_empresa_planes.precio_plan":"Precio del plan anterior.",
 "historial_empresa_planes.cantidad_reportes":"Reportes del plan anterior.","historial_empresa_planes.precio_sobre_fact":"Precio de reporte extra del plan anterior.","historial_empresa_planes.fecha_desde":"Fecha de inicio de vigencia.","historial_empresa_planes.fecha_hasta":"Fecha de fin de vigencia.",
 "historial_empresa_planes.cambiado_por":"Usuario que realizo el cambio de plan.","historial_empresa_planes.fecha_cambio":"Fecha del cambio de plan.",
 "auditoria.id":"Identificador unico del evento.","auditoria.usuario_id":"Usuario que ejecuto la accion.","auditoria.usuario_nombre":"Nombre del usuario (copia al momento del evento).","auditoria.id_empresa":"Empresa asociada al evento.",
 "auditoria.tipo_evento":"Tipo: AUTENTICACION, CLIENTES, EVALUACIONES, etc.","auditoria.accion":"Accion realizada (ej: LOGIN_EXITOSO, REGISTRO_CLIENTE).","auditoria.modulo":"Modulo del sistema donde ocurrio el evento.","auditoria.entidad":"Entidad afectada (ej: Cliente, Usuario).",
 "auditoria.registro_id":"Identificador del registro afectado.","auditoria.resultado":"Resultado del evento: EXITO o FALLO.","auditoria.ip":"Direccion IP desde donde se realizo la accion.","auditoria.info_adicional":"Informacion adicional en formato JSON.",
 "auditoria.valores_anteriores":"Valores antes de la modificacion (JSON).","auditoria.valores_nuevos":"Valores despues de la modificacion (JSON).","auditoria.fecha":"Fecha y hora del evento.",
 "lista_negra_onu.id":"Identificador unico.","lista_negra_onu.registro":"Codigo de referencia del registro.","lista_negra_onu.nombre":"Nombre de la persona/entidad sancionada.","lista_negra_onu.apellido":"Apellido de la persona sancionada.",
 "lista_negra_onu.cargo":"Cargo o descripcion.","lista_negra_onu.fecha_nacimiento":"Fecha de nacimiento (texto).","lista_negra_onu.nacionalidad":"Nacionalidad.","lista_negra_onu.num_identidad":"Numero de identidad nacional.","lista_negra_onu.num_pasaporte":"Numero de pasaporte.","lista_negra_onu.otros":"Informacion adicional.",
}

ORDEN = ["usuarios","empresas","clientes","cliente_empresa","paises","departamento","ciudad",
 "documentos","historial_crediticio","criterios","evaluaciones","evaluacion_detalle",
 "motor_reglas","reglas","evaluacion_riesgo","resultado_detalle","detalle_operacion",
 "scoring_modelo","scoring_factor","scoring_catalogo","scoring_regla","scoring_umbral",
 "scoring_evaluacion","scoring_detalle","planes","empresa_planes","periodos_facturacion",
 "consumo_reportes","pagos","historial_empresa_planes","auditoria","lista_negra_onu"]


def longitud(tipo, ml, pr, sc):
    if tipo in ("varchar","char"): return "MAX" if ml==-1 else str(ml)
    if tipo=="nvarchar": return "MAX" if ml==-1 else str(ml//2)
    if tipo in ("decimal","numeric"): return f"{pr},{sc}"
    if tipo=="bit": return "1"
    return "-"


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


app = create_app()
with app.app_context():
    conn = db.engine.connect()
    pks = conn.execute(text("""
        SELECT t.name tabla, c.name col FROM sys.indexes i
        JOIN sys.index_columns ic ON i.object_id=ic.object_id AND i.index_id=ic.index_id
        JOIN sys.columns c ON ic.object_id=c.object_id AND ic.column_id=c.column_id
        JOIN sys.tables t ON i.object_id=t.object_id WHERE i.is_primary_key=1
    """))
    pk_set = {(r[0],r[1]) for r in pks}
    fks = conn.execute(text("""
        SELECT tr.name t, cr.name c, tp.name rt FROM sys.foreign_keys fk
        JOIN sys.foreign_key_columns fkc ON fk.object_id=fkc.constraint_object_id
        JOIN sys.tables tp ON fkc.referenced_object_id=tp.object_id
        JOIN sys.tables tr ON fkc.parent_object_id=tr.object_id
        JOIN sys.columns cr ON fkc.parent_object_id=cr.object_id AND fkc.parent_column_id=cr.column_id
    """))
    fk_map = {(r[0],r[1]): r[2] for r in fks}
    cols = conn.execute(text("""
        SELECT t.name tabla, c.name col, ty.name tipo, c.max_length ml, c.precision pr,
               c.scale sc, c.is_nullable nul, c.is_identity idd
        FROM sys.tables t JOIN sys.columns c ON t.object_id=c.object_id
        JOIN sys.types ty ON c.user_type_id=ty.user_type_id
        WHERE t.name NOT IN ('sysdiagrams') ORDER BY t.name, c.column_id
    """))
    tablas = {}
    for r in cols:
        tablas.setdefault(r[0], []).append(r)

    doc = Document()
    # margenes
    for s in doc.sections:
        s.left_margin = Cm(2); s.right_margin = Cm(2)
        s.top_margin = Cm(2); s.bottom_margin = Cm(2)

    # Titulo
    h = doc.add_heading("Diccionario de Datos", level=0)
    p = doc.add_paragraph("Sistema Experto de Scoring Comercial Inmobiliario")
    p.runs[0].italic = True
    doc.add_paragraph("Base de datos: inmobiliaria_db  |  Motor: Microsoft SQL Server 2022")

    orden_final = [t for t in ORDEN if t in tablas] + [t for t in tablas if t not in ORDEN]

    for tabla in orden_final:
        doc.add_heading(f"Tabla: {tabla}", level=2)
        if DESC_TABLA.get(tabla):
            pd = doc.add_paragraph()
            r = pd.add_run("Descripcion: ")
            r.bold = True
            pd.add_run(DESC_TABLA[tabla])

        tbl = doc.add_table(rows=1, cols=4)
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        # anchos
        widths = [Cm(3.5), Cm(3.0), Cm(1.6), Cm(9.0)]
        hdr = tbl.rows[0].cells
        encabezados = ["Campo","Tipo de Dato","Long.","Descripcion / Restricciones"]
        for i, txt in enumerate(encabezados):
            hdr[i].text = ""
            pr = hdr[i].paragraphs[0]
            run = pr.add_run(txt)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
            run.font.size = Pt(9)
            set_cell_bg(hdr[i], "1E3A5F")
            hdr[i].width = widths[i]

        for rr in tablas[tabla]:
            _, col, tipo, ml, pr_, sc, nul, idd = rr
            tt = tipo.upper()
            if (tabla,col) in pk_set: tt += " (PK)"
            elif (tabla,col) in fk_map: tt += " (FK)"
            base = DESC_CAMPO.get(f"{tabla}.{col}", "")
            restr = []
            if (tabla,col) in fk_map: restr.append(f"FK -> {fk_map[(tabla,col)]}")
            if idd: restr.append("Autoincremental")
            if not nul: restr.append("Obligatorio")
            texto = base
            if restr:
                texto = (base + " " if base else "") + "(" + "; ".join(restr) + ")"
            row = tbl.add_row().cells
            valores = [col, tt, longitud(tipo,ml,pr_,sc), texto or "-"]
            for i, v in enumerate(valores):
                row[i].text = ""
                run = row[i].paragraphs[0].add_run(v)
                run.font.size = Pt(9)
                row[i].width = widths[i]

        doc.add_paragraph("")

    doc.save(OUT)
    print("WORD generado:", OUT)
