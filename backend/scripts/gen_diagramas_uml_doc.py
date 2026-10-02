# -*- coding: utf-8 -*-
"""
Genera un Word con el código PlantUML de los diagramas de Actividades, Clases,
Componentes y Despliegue, listo para pegar en https://editor.plantuml.com/uml
"""
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SALIDA = r"C:\Users\RYZEN\Documents\TESIS 2026_DOCUMENTOS\Diagramas_UML_PlantUML.docx"
BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "diagramas"))

DIAGRAMAS = [
    ("Diagrama de Actividades — Evaluar Riesgo de Cliente",
     "Modela el flujo de trabajo del caso de uso central del sistema experto, "
     "desde la selección del cliente hasta la presentación del resultado, "
     "incluyendo la verificación de lista negra, la evaluación en paralelo de "
     "las reglas, las decisiones de clasificación del riesgo y el control de "
     "consumo.",
     os.path.join(BASE, "Actividades_Evaluar_Riesgo.puml")),
    ("Diagrama de Clases — Modelo de Dominio",
     "Representa la estructura estática del sistema: las entidades principales "
     "con sus atributos y métodos, y las relaciones de asociación, agregación y "
     "composición entre ellas.",
     os.path.join(BASE, "Clases_Sistema.puml")),
    ("Diagrama de Componentes — Arquitectura del Software",
     "Representa los componentes lógicos del sistema (frontend React, módulos "
     "del backend Flask, núcleo de scoring, servicios y capa ORM) y sus "
     "relaciones e interfaces.",
     os.path.join(BASE, "Componentes_Sistema.puml")),
    ("Diagrama de Despliegue — Arquitectura Física",
     "Representa la distribución física del sistema: el navegador del usuario, "
     "el sitio estático y el servicio web desplegados en Render, la base de "
     "datos PostgreSQL administrada y la pasarela de pagos externa.",
     os.path.join(BASE, "Despliegue_Sistema.puml")),
]


def add_code_block(doc, texto):
    for linea in texto.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:ascii'), 'Consolas')
        rfonts.set(qn('w:hAnsi'), 'Consolas')


def main():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)

    doc.add_heading('Diagramas UML — Actividades, Clases, Componentes y Despliegue', level=1)
    p = doc.add_paragraph(
        'A continuación se presenta el código fuente en PlantUML de los diagramas '
        'de actividades, clases, componentes y despliegue del sistema. Para generar '
        'la imagen de cada diagrama, copie el bloque de código correspondiente y '
        'péguelo en el editor en línea https://editor.plantuml.com/uml'
    )
    p.paragraph_format.line_spacing = 1.5

    for titulo, desc, ruta in DIAGRAMAS:
        doc.add_paragraph()
        doc.add_heading(titulo, level=2)
        d = doc.add_paragraph()
        r = d.add_run('Descripción: ')
        r.bold = True
        d.add_run(desc)
        d.paragraph_format.line_spacing = 1.5
        doc.add_paragraph()
        try:
            with open(ruta, "r", encoding="utf-8") as fh:
                contenido = fh.read()
        except FileNotFoundError:
            contenido = f"' (No se encontró el archivo: {ruta})"
        add_code_block(doc, contenido)

    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    doc.save(SALIDA)
    print("Documento generado:", SALIDA)


if __name__ == "__main__":
    main()
