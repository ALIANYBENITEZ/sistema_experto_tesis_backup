# -*- coding: utf-8 -*-
"""
Genera un Word con el código PlantUML de los diagramas de secuencia,
listo para pegar en https://editor.plantuml.com/uml
"""
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SALIDA = r"C:\Users\RYZEN\Documents\TESIS 2026_DOCUMENTOS\Diagramas_Secuencia_PlantUML.docx"
BASE = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "diagramas"))

DIAGRAMAS = [
    ("Diagrama de Secuencia — Inicio de Sesión con Doble Factor (2FA)",
     "Representa el flujo del caso de uso de autenticación: validación de "
     "credenciales, verificación del segundo factor (código TOTP de Google "
     "Authenticator) y emisión de los tokens de acceso. Incluye tanto el "
     "escenario de 2FA ya configurado como la configuración inicial.",
     os.path.join(BASE, "Secuencia_Login_2FA.puml")),
    ("Diagrama de Secuencia — Registrar Empresa y Usuario",
     "Representa la gestión, por parte del propietario del sistema, del alta de "
     "una empresa cliente y la creación de los usuarios (administrador o "
     "comercial) que pertenecen a esa empresa.",
     os.path.join(BASE, "Secuencia_Registrar_Empresa_Usuario.puml")),
    ("Diagrama de Secuencia — Registrar Cliente",
     "Representa el registro de un cliente por parte del operador, con las "
     "validaciones de mayoría de edad y teléfono, la detección de clientes "
     "existentes y la vinculación del cliente a la empresa.",
     os.path.join(BASE, "Secuencia_Registrar_Cliente.puml")),
    ("Diagrama de Secuencia — Configurar Motor de Reglas",
     "Representa la definición del motor de reglas IF-THEN del sistema experto: "
     "creación del motor y alta de las reglas (parámetro, operador, valor, peso "
     "y regla determinante).",
     os.path.join(BASE, "Secuencia_Configurar_Motor_Reglas.puml")),
    ("Diagrama de Secuencia — Evaluar Riesgo de Cliente",
     "Representa el caso de uso central del sistema experto: ejecución del "
     "motor de scoring sobre un cliente, verificación contra la lista de "
     "sanciones, evaluación de las reglas IF-THEN, clasificación del riesgo, "
     "registro del consumo y auditoría del evento.",
     os.path.join(BASE, "Secuencia_Evaluar_Riesgo.puml")),
    ("Diagrama de Secuencia — Gestionar Pago de Facturación",
     "Representa el pago de un período de facturación mediante la pasarela "
     "Pagopar, incluyendo la notificación asincrónica (webhook) y el desbloqueo "
     "automático de la empresa al saldar la deuda.",
     os.path.join(BASE, "Secuencia_Gestionar_Pago.puml")),
    ("Diagrama de Secuencia — Generar Reporte PDF",
     "Representa la exportación del reporte de evaluaciones de riesgo a PDF, con "
     "aplicación de filtros y control de visibilidad según el rol del usuario.",
     os.path.join(BASE, "Secuencia_Generar_Reporte_PDF.puml")),
    ("Diagrama de Secuencia — Consultar Auditoría",
     "Representa la consulta de la bitácora de eventos del sistema, con filtros "
     "y paginación, respetando el alcance de visibilidad según el rol.",
     os.path.join(BASE, "Secuencia_Consultar_Auditoria.puml")),
]


def add_code_block(doc, texto):
    for linea in texto.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.append(rfonts)
        rfonts.set(qn('w:ascii'), 'Consolas')
        rfonts.set(qn('w:hAnsi'), 'Consolas')


def main():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)

    doc.add_heading('Diagramas de Secuencia (PlantUML)', level=1)
    p = doc.add_paragraph(
        'A continuación se presenta el código fuente en PlantUML de los diagramas '
        'de secuencia del sistema. Para generar la imagen del diagrama, copie el '
        'bloque de código correspondiente y péguelo en el editor en línea '
        'https://editor.plantuml.com/uml'
    )
    p.paragraph_format.line_spacing = 1.5

    for titulo, desc, ruta in DIAGRAMAS:
        doc.add_paragraph()
        doc.add_heading(titulo, level=2)
        d = doc.add_paragraph()
        r = d.add_run('Descripción: ')
        r.bold = True
        d.add_run(desc)
        d.paragraph_format.line_spacing = 1.5
        doc.add_paragraph()
        try:
            with open(ruta, "r", encoding="utf-8") as fh:
                contenido = fh.read()
        except FileNotFoundError:
            contenido = f"' (No se encontró el archivo: {ruta})"
        add_code_block(doc, contenido)

    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    doc.save(SALIDA)
    print("Documento generado:", SALIDA)


if __name__ == "__main__":
    main()
