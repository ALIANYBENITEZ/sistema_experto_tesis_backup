# -*- coding: utf-8 -*-
"""
Genera el Diccionario de Datos (Word .docx) del sistema, basado en la base
de datos PostgreSQL real. Estructura por tabla:
  Encabezado: nombre + descripción
  Tabla: Campo | Tipo | Long. | Descripción / Restricciones
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SALIDA = r"C:\Users\RYZEN\Documents\TESIS 2026_DOCUMENTOS\Diccionario_de_Datos.docx"

# Cada tabla: (nombre, descripción, [ (campo, tipo, long, descripcion_restricciones), ... ])
TABLAS = [
    ("usuarios", "Usuarios del sistema con sus credenciales, rol y configuración de segundo factor (2FA).", [
        ("id", "INTEGER (PK)", "-", "Identificador único del usuario. (Autoincremental; Obligatorio)"),
        ("nombre", "VARCHAR", "100", "Nombre del usuario. (Obligatorio)"),
        ("apellido", "VARCHAR", "100", "Apellido del usuario. (Obligatorio)"),
        ("email", "VARCHAR", "150", "Correo electrónico. (Obligatorio; Único)"),
        ("password_hash", "VARCHAR", "255", "Contraseña cifrada con scrypt. (Obligatorio)"),
        ("rol", "VARCHAR", "20", "Rol del usuario: propietario, administrador, comercial. (Obligatorio)"),
        ("activo", "BOOLEAN", "-", "Indica si el usuario está habilitado. (Por defecto: verdadero)"),
        ("id_empresa", "INTEGER", "-", "Empresa a la que pertenece (0 = sistema/propietario)."),
        ("totp_secret", "VARCHAR", "255", "Secreto TOTP para 2FA. (Opcional)"),
        ("totp_estado", "VARCHAR", "20", "Estado del 2FA: no_configurado, pendiente, activado, restablecido."),
        ("totp_fecha_config", "TIMESTAMP", "-", "Fecha de configuración del 2FA. (Opcional)"),
        ("totp_fecha_reset", "TIMESTAMP", "-", "Fecha de restablecimiento del 2FA. (Opcional)"),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación del registro."),
        ("actualizado_en", "TIMESTAMP", "-", "Fecha de última actualización."),
    ]),
    ("empresas", "Empresas clientes que contratan el uso del sistema.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la empresa. (Autoincremental; Obligatorio)"),
        ("nombre", "VARCHAR", "150", "Razón social o nombre de la empresa. (Obligatorio)"),
        ("ruc", "VARCHAR", "30", "RUC de la empresa. (Único)"),
        ("direccion", "VARCHAR", "255", "Dirección de la empresa. (Opcional)"),
        ("telefono", "VARCHAR", "30", "Teléfono de contacto. (Opcional)"),
        ("email", "VARCHAR", "150", "Correo de contacto. (Opcional)"),
        ("activo", "BOOLEAN", "-", "Indica si la empresa está activa. (Por defecto: verdadero)"),
        ("consultas_habilitadas", "BOOLEAN", "-", "Habilita o bloquea el consumo de reportes. (Por defecto: verdadero)"),
        ("motivo_bloqueo", "VARCHAR", "50", "Motivo del bloqueo: BLOQUEADO_DEUDA, BLOQUEADO_MANUAL. (Opcional)"),
        ("fecha_bloqueo", "TIMESTAMP", "-", "Fecha en que se bloqueó. (Opcional)"),
        ("fecha_desbloqueo", "TIMESTAMP", "-", "Fecha en que se desbloqueó. (Opcional)"),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación del registro."),
    ]),
    ("paises", "Catálogo de países usado para la nacionalidad de clientes.", [
        ("id_pais", "VARCHAR (PK)", "10", "Código del país (ej. PY, AR, BR). (Obligatorio)"),
        ("nombre_pais", "VARCHAR", "100", "Nombre del país. (Obligatorio)"),
    ]),
    ("departamento", "Departamentos de Paraguay.", [
        ("id_departamento", "INTEGER (PK)", "-", "Identificador único del departamento. (Obligatorio)"),
        ("nombre_departamento", "VARCHAR", "100", "Nombre del departamento. (Obligatorio)"),
    ]),
    ("ciudad", "Ciudades asociadas a un departamento.", [
        ("id_ciudad", "INTEGER (PK)", "-", "Identificador único de la ciudad. (Obligatorio)"),
        ("nombre_ciudad", "VARCHAR", "100", "Nombre de la ciudad. (Obligatorio)"),
        ("id_departamento", "INTEGER (FK)", "-", "Departamento al que pertenece. (Obligatorio; ref. departamento)"),
    ]),
    ("clientes", "Clientes que solicitan evaluación para adquirir una propiedad.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del cliente. (Autoincremental; Obligatorio)"),
        ("tipo_doc", "VARCHAR", "10", "Tipo de documento: CI, RUC, PAS. (Obligatorio)"),
        ("num_doc", "VARCHAR", "20", "Número de documento. (Obligatorio; Único)"),
        ("nombre", "VARCHAR", "100", "Nombre del cliente. (Obligatorio)"),
        ("apellido", "VARCHAR", "100", "Apellido del cliente. (Opcional)"),
        ("email", "VARCHAR", "150", "Correo electrónico. (Opcional)"),
        ("telefono", "VARCHAR", "20", "Teléfono (solo dígitos y símbolos + - ( ) espacio). (Opcional)"),
        ("direccion", "VARCHAR", "255", "Dirección del cliente. (Opcional)"),
        ("fecha_nacimiento", "DATE", "-", "Fecha de nacimiento (el cliente debe ser mayor de edad). (Opcional)"),
        ("nacionalidad", "VARCHAR (FK)", "10", "País de nacionalidad. (Opcional; ref. paises)"),
        ("id_ciudad", "INTEGER (FK)", "-", "Ciudad del cliente. (Opcional; ref. ciudad)"),
        ("estado", "VARCHAR", "20", "Estado del cliente: activo, inactivo. (Por defecto: activo)"),
        ("creado_por", "INTEGER (FK)", "-", "Usuario que registró al cliente. (ref. usuarios)"),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación del registro."),
    ]),
    ("cliente_empresa", "Relación muchos a muchos entre clientes y empresas.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del vínculo. (Autoincremental; Obligatorio)"),
        ("id_cliente", "INTEGER (FK)", "-", "Cliente vinculado. (Obligatorio; ref. clientes)"),
        ("id_empresa", "INTEGER (FK)", "-", "Empresa vinculada. (Obligatorio; ref. empresas)"),
        ("estado", "VARCHAR", "20", "Estado del vínculo: activo, inactivo. (Por defecto: activo)"),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación del vínculo."),
        ("(uq)", "UNIQUE", "-", "Restricción única sobre (id_cliente, id_empresa)."),
    ]),
    ("criterios", "Criterios de evaluación ponderada configurables por el administrador.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del criterio. (Autoincremental; Obligatorio)"),
        ("nombre", "VARCHAR", "100", "Nombre del criterio. (Obligatorio)"),
        ("descripcion", "VARCHAR", "255", "Descripción del criterio. (Opcional)"),
        ("peso", "NUMERIC", "5,2", "Peso porcentual (0 < peso <= 100). (Obligatorio)"),
        ("activo", "BOOLEAN", "-", "Indica si el criterio está activo. (Por defecto: verdadero)"),
    ]),
    ("evaluaciones", "Registro de cada evaluación por criterios realizada a un cliente.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la evaluación. (Autoincremental; Obligatorio)"),
        ("cliente_id", "INTEGER (FK)", "-", "Cliente evaluado. (Obligatorio; ref. clientes)"),
        ("evaluador_id", "INTEGER (FK)", "-", "Usuario que realizó la evaluación. (Obligatorio; ref. usuarios)"),
        ("fecha", "TIMESTAMP", "-", "Fecha de la evaluación."),
        ("puntaje_total", "NUMERIC", "5,2", "Puntaje total obtenido. (Opcional)"),
        ("resultado", "VARCHAR", "20", "Resultado: aprobado, observado, rechazado. (Opcional)"),
        ("observaciones", "TEXT", "-", "Observaciones de la evaluación. (Opcional)"),
        ("estado", "VARCHAR", "20", "Estado: pendiente, completado, anulado. (Por defecto: pendiente)"),
    ]),
    ("evaluacion_detalle", "Detalle por criterio de cada evaluación.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del detalle. (Autoincremental; Obligatorio)"),
        ("evaluacion_id", "INTEGER (FK)", "-", "Evaluación asociada. (Obligatorio; ref. evaluaciones; ON DELETE CASCADE)"),
        ("criterio_id", "INTEGER (FK)", "-", "Criterio evaluado. (Obligatorio; ref. criterios)"),
        ("valor", "NUMERIC", "5,2", "Valor asignado (0 a 100). (Obligatorio)"),
        ("comentario", "VARCHAR", "255", "Comentario del criterio. (Opcional)"),
    ]),
    ("documentos", "Documentos adjuntos de los clientes (DNI, recibos, etc.).", [
        ("id", "INTEGER (PK)", "-", "Identificador único del documento. (Autoincremental; Obligatorio)"),
        ("cliente_id", "INTEGER (FK)", "-", "Cliente propietario del documento. (Obligatorio; ref. clientes)"),
        ("tipo", "VARCHAR", "50", "Tipo de documento. (Obligatorio)"),
        ("nombre_archivo", "VARCHAR", "255", "Nombre del archivo. (Obligatorio)"),
        ("ruta", "VARCHAR", "500", "Ruta de almacenamiento. (Obligatorio)"),
        ("subido_por", "INTEGER (FK)", "-", "Usuario que subió el documento. (ref. usuarios)"),
        ("subido_en", "TIMESTAMP", "-", "Fecha de subida."),
    ]),
    ("motor_reglas", "Plantilla (motor) de reglas IF-THEN asignable a una empresa.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del motor. (Autoincremental; Obligatorio)"),
        ("nombre", "VARCHAR", "100", "Nombre del motor de reglas. (Obligatorio)"),
        ("version", "VARCHAR", "20", "Versión del motor. (Por defecto: 1.0)"),
        ("descripcion", "VARCHAR", "255", "Descripción del motor. (Opcional)"),
        ("activo", "BOOLEAN", "-", "Indica si el motor está activo. (Por defecto: verdadero)"),
        ("creado_por", "INTEGER (FK)", "-", "Usuario creador. (ref. usuarios)"),
        ("id_empresa_motor", "INTEGER", "-", "Empresa dueña del motor (0 = global)."),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación."),
    ]),
    ("reglas", "Regla individual de evaluación (condición IF-THEN) de un motor.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la regla. (Autoincremental; Obligatorio)"),
        ("motor_id", "INTEGER (FK)", "-", "Motor al que pertenece. (Obligatorio; ref. motor_reglas; ON DELETE CASCADE)"),
        ("nombre", "VARCHAR", "100", "Nombre de la regla. (Obligatorio)"),
        ("descripcion", "VARCHAR", "255", "Descripción de la regla. (Opcional)"),
        ("parametro", "VARCHAR", "50", "Parámetro a evaluar (ej. ingresos_mensuales). (Obligatorio)"),
        ("operador", "VARCHAR", "10", "Operador lógico: >=, <=, ==, >, <, !=. (Obligatorio)"),
        ("valor_referencia", "VARCHAR", "100", "Umbral de comparación. (Obligatorio)"),
        ("tipo_valor", "VARCHAR", "20", "Tipo de valor: numerico, texto, booleano. (Por defecto: numerico)"),
        ("peso_puntos", "NUMERIC", "6,2", "Puntos otorgados si se cumple. (Obligatorio)"),
        ("es_determinante", "BOOLEAN", "-", "Si falla, produce rechazo directo. (Por defecto: falso)"),
        ("activo", "BOOLEAN", "-", "Indica si la regla está activa. (Por defecto: verdadero)"),
        ("orden", "INTEGER", "-", "Orden de evaluación. (Por defecto: 0)"),
    ]),
    ("historial_crediticio", "Datos financieros e histórico del cliente para la evaluación de riesgo.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del registro. (Autoincremental; Obligatorio)"),
        ("cliente_id", "INTEGER (FK)", "-", "Cliente asociado. (Obligatorio; Único; ref. clientes)"),
        ("fuente_externa", "VARCHAR", "100", "Fuente de los datos externos. (Opcional)"),
        ("score_externo", "NUMERIC", "6,2", "Score de central de riesgo. (Opcional)"),
        ("cantidad_atrasos", "INTEGER", "-", "Cantidad de atrasos. (Por defecto: 0)"),
        ("deuda_total_sistema", "NUMERIC", "14,2", "Deuda total en el sistema financiero. (Por defecto: 0)"),
        ("historial_pagos", "VARCHAR", "20", "bueno, regular, malo, sin_historial. (Por defecto: sin_historial)"),
        ("en_lista_negra", "BOOLEAN", "-", "Si está en lista OFAC/ONU. (Por defecto: falso)"),
        ("nivel_endeudamiento", "NUMERIC", "5,2", "Porcentaje de endeudamiento. (Por defecto: 0)"),
        ("meses_empleo_actual", "INTEGER", "-", "Antigüedad laboral en meses. (Por defecto: 0)"),
        ("tipo_empleo", "VARCHAR", "50", "dependiente, independiente, desempleado. (Opcional)"),
        ("referencias_personales", "VARCHAR", "20", "buenas, regulares, malas, no_verificadas. (Por defecto: no_verificadas)"),
        ("fecha_consulta", "TIMESTAMP", "-", "Fecha de consulta de los datos."),
        ("actualizado_en", "TIMESTAMP", "-", "Fecha de última actualización."),
    ]),
    ("evaluacion_riesgo", "Resultado de una evaluación de riesgo (motor de reglas) de un cliente.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la evaluación. (Autoincremental; Obligatorio)"),
        ("cliente_id", "INTEGER (FK)", "-", "Cliente evaluado. (Obligatorio; ref. clientes)"),
        ("motor_id", "INTEGER (FK)", "-", "Motor de reglas aplicado. (Obligatorio; ref. motor_reglas)"),
        ("usuario_id", "INTEGER (FK)", "-", "Usuario que ejecutó la evaluación. (Obligatorio; ref. usuarios)"),
        ("fecha_analisis", "TIMESTAMP", "-", "Fecha del análisis."),
        ("score_final", "NUMERIC", "6,2", "Puntaje final obtenido. (Opcional)"),
        ("score_maximo", "NUMERIC", "6,2", "Puntaje máximo posible. (Opcional)"),
        ("categoria_riesgo", "VARCHAR", "20", "bajo, medio, alto, rechazado. (Opcional)"),
        ("estado", "VARCHAR", "20", "Estado de la evaluación. (Por defecto: completado)"),
        ("observaciones", "TEXT", "-", "Observaciones del análisis. (Opcional)"),
        ("regla_determinante_id", "INTEGER (FK)", "-", "Regla determinante que causó rechazo. (Opcional; ref. reglas)"),
        ("id_empresa", "INTEGER", "-", "Empresa asociada a la evaluación."),
    ]),
    ("resultado_detalle", "Detalle por regla de una evaluación de riesgo.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del detalle. (Autoincremental; Obligatorio)"),
        ("evaluacion_id", "INTEGER (FK)", "-", "Evaluación de riesgo asociada. (Obligatorio; ref. evaluacion_riesgo; ON DELETE CASCADE)"),
        ("regla_id", "INTEGER (FK)", "-", "Regla evaluada. (Obligatorio; ref. reglas)"),
        ("cumplido", "BOOLEAN", "-", "Indica si la regla se cumplió. (Obligatorio)"),
        ("valor_evaluado", "VARCHAR", "100", "Valor real del cliente al momento. (Opcional)"),
        ("puntos_obtenidos", "NUMERIC", "6,2", "Puntos obtenidos por la regla. (Por defecto: 0)"),
    ]),
    ("detalle_operacion", "Datos de la operación inmobiliaria asociada a una evaluación de riesgo.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la operación. (Autoincremental; Obligatorio)"),
        ("evaluacion_id", "INTEGER (FK)", "-", "Evaluación de riesgo asociada. (Obligatorio; ref. evaluacion_riesgo; ON DELETE CASCADE)"),
        ("tipo_propiedad", "VARCHAR", "50", "casa, departamento, terreno, local_comercial, oficina. (Obligatorio)"),
        ("valor_propiedad", "NUMERIC", "14,2", "Valor de la propiedad. (Obligatorio)"),
        ("monto_solicitado", "NUMERIC", "14,2", "Monto solicitado. (Obligatorio)"),
        ("plazo_meses", "INTEGER", "-", "Plazo en meses. (Obligatorio)"),
        ("ubicacion", "VARCHAR", "255", "Ubicación de la propiedad. (Opcional)"),
        ("destino", "VARCHAR", "50", "vivienda, inversion, comercial. (Opcional)"),
    ]),
    ("lista_negra_onu", "Lista consolidada de sanciones de la ONU (referencia externa para verificación).", [
        ("id", "INTEGER (PK)", "-", "Identificador único del registro. (Autoincremental; Obligatorio)"),
        ("registro", "VARCHAR", "20", "Código de referencia (ej. CDi.001). (Opcional)"),
        ("nombre", "VARCHAR", "150", "Nombre. (Opcional)"),
        ("apellido", "VARCHAR", "150", "Apellido. (Opcional)"),
        ("cargo", "VARCHAR", "500", "Cargo. (Opcional)"),
        ("fecha_nacimiento", "VARCHAR", "200", "Fecha(s) de nacimiento. (Opcional)"),
        ("nacionalidad", "VARCHAR", "200", "Nacionalidad. (Opcional)"),
        ("num_identidad", "VARCHAR", "300", "Número(s) de identidad. (Opcional)"),
        ("num_pasaporte", "VARCHAR", "500", "Número(s) de pasaporte. (Opcional)"),
        ("otros", "TEXT", "-", "Otros datos. (Opcional)"),
    ]),
    ("planes", "Catálogo de planes comerciales (prepago por reportes).", [
        ("id", "INTEGER (PK)", "-", "Identificador único del plan. (Autoincremental; Obligatorio)"),
        ("nombre", "VARCHAR", "50", "Nombre del plan. (Obligatorio; Único)"),
        ("cantidad_reportes_incluidos", "INTEGER", "-", "Reportes incluidos en el plan. (Obligatorio)"),
        ("precio_plan", "NUMERIC", "14,0", "Precio del plan en Gs. (Obligatorio)"),
        ("precio_reporte_incluido", "NUMERIC", "14,0", "Precio de referencia por reporte incluido. (Obligatorio)"),
        ("precio_reporte_sobre_facturado", "NUMERIC", "14,0", "Precio por reporte excedente. (Obligatorio)"),
        ("activo", "BOOLEAN", "-", "Indica si el plan está activo. (Por defecto: verdadero)"),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación."),
        ("actualizado_en", "TIMESTAMP", "-", "Fecha de última actualización."),
    ]),
    ("empresa_planes", "Plan vigente asignado a una empresa (con copia de valores al contratar).", [
        ("id", "INTEGER (PK)", "-", "Identificador único. (Autoincremental; Obligatorio)"),
        ("empresa_id", "INTEGER (FK)", "-", "Empresa. (Obligatorio; ref. empresas)"),
        ("plan_id", "INTEGER (FK)", "-", "Plan contratado. (Obligatorio; ref. planes)"),
        ("fecha_inicio", "TIMESTAMP", "-", "Inicio de vigencia. (Obligatorio)"),
        ("fecha_fin", "TIMESTAMP", "-", "Fin de vigencia. (Opcional)"),
        ("estado", "VARCHAR", "20", "activo, finalizado. (Por defecto: activo)"),
        ("precio_plan_contratado", "NUMERIC", "14,0", "Precio del plan al contratar. (Obligatorio)"),
        ("cantidad_reportes_incluidos", "INTEGER", "-", "Reportes incluidos al contratar. (Obligatorio)"),
        ("precio_reporte_sobre_facturado", "NUMERIC", "14,0", "Precio de excedente al contratar. (Obligatorio)"),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación."),
    ]),
    ("periodos_facturacion", "Período mensual de facturación por empresa.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del período. (Autoincremental; Obligatorio)"),
        ("empresa_id", "INTEGER (FK)", "-", "Empresa. (Obligatorio; ref. empresas)"),
        ("empresa_plan_id", "INTEGER (FK)", "-", "Plan de empresa vigente. (Obligatorio; ref. empresa_planes)"),
        ("anio", "INTEGER", "-", "Año del período. (Obligatorio)"),
        ("mes", "INTEGER", "-", "Mes del período. (Obligatorio)"),
        ("fecha_inicio", "TIMESTAMP", "-", "Inicio del período. (Obligatorio)"),
        ("fecha_fin", "TIMESTAMP", "-", "Fin del período. (Obligatorio)"),
        ("cantidad_reportes_incluidos", "INTEGER", "-", "Reportes incluidos en el período. (Obligatorio)"),
        ("reportes_consumidos", "INTEGER", "-", "Reportes consumidos. (Por defecto: 0)"),
        ("reportes_sobre_facturados", "INTEGER", "-", "Reportes excedentes. (Por defecto: 0)"),
        ("monto_plan", "NUMERIC", "14,0", "Monto del plan. (Por defecto: 0)"),
        ("monto_sobre_facturado", "NUMERIC", "14,0", "Monto por excedentes. (Por defecto: 0)"),
        ("monto_total", "NUMERIC", "14,0", "Monto total del período. (Por defecto: 0)"),
        ("monto_pagado", "NUMERIC", "14,0", "Monto pagado. (Por defecto: 0)"),
        ("saldo_pendiente", "NUMERIC", "14,0", "Saldo pendiente. (Por defecto: 0)"),
        ("estado", "VARCHAR", "20", "PENDIENTE, PARCIAL, PAGADO, VENCIDO, BLOQUEADO. (Por defecto: PENDIENTE)"),
        ("(uq)", "UNIQUE", "-", "Restricción única sobre (empresa_id, anio, mes)."),
    ]),
    ("consumo_reportes", "Registro individual de cada reporte consumido.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del consumo. (Autoincremental; Obligatorio)"),
        ("empresa_id", "INTEGER (FK)", "-", "Empresa que consume. (Obligatorio; ref. empresas)"),
        ("usuario_id", "INTEGER", "-", "Usuario que generó el reporte. (Obligatorio)"),
        ("evaluacion_id", "INTEGER", "-", "Evaluación asociada. (Obligatorio)"),
        ("periodo_facturacion_id", "INTEGER (FK)", "-", "Período de facturación. (Obligatorio; ref. periodos_facturacion)"),
        ("tipo_consumo", "VARCHAR", "20", "INCLUIDO, SOBRE_FACTURADO. (Obligatorio)"),
        ("precio_unitario", "NUMERIC", "14,0", "Precio unitario aplicado. (Obligatorio)"),
        ("fecha_consumo", "TIMESTAMP", "-", "Fecha del consumo."),
        ("(uq)", "UNIQUE", "-", "Restricción única sobre (evaluacion_id, empresa_id)."),
    ]),
    ("pagos", "Transacciones de pago de las empresas.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del pago. (Autoincremental; Obligatorio)"),
        ("empresa_id", "INTEGER (FK)", "-", "Empresa que paga. (Obligatorio; ref. empresas)"),
        ("periodo_facturacion_id", "INTEGER (FK)", "-", "Período pagado. (Obligatorio; ref. periodos_facturacion)"),
        ("monto", "NUMERIC", "14,0", "Monto del pago. (Obligatorio)"),
        ("proveedor", "VARCHAR", "50", "TEST, BANCARD, PAGOPAR. (Por defecto: TEST)"),
        ("referencia_externa", "VARCHAR", "100", "Referencia de la pasarela. (Opcional)"),
        ("estado", "VARCHAR", "20", "PENDIENTE, APROBADO, RECHAZADO, CANCELADO. (Por defecto: PENDIENTE)"),
        ("metodo_pago", "VARCHAR", "50", "Método de pago. (Opcional)"),
        ("observacion", "VARCHAR", "255", "Observación. (Opcional)"),
        ("fecha_inicio", "TIMESTAMP", "-", "Fecha de inicio del pago."),
        ("fecha_confirmacion", "TIMESTAMP", "-", "Fecha de confirmación. (Opcional)"),
    ]),
    ("historial_empresa_planes", "Historial de cambios de plan por empresa.", [
        ("id", "INTEGER (PK)", "-", "Identificador único. (Autoincremental; Obligatorio)"),
        ("empresa_id", "INTEGER (FK)", "-", "Empresa. (Obligatorio; ref. empresas)"),
        ("plan_id", "INTEGER (FK)", "-", "Plan. (Obligatorio; ref. planes)"),
        ("plan_nombre", "VARCHAR", "50", "Nombre del plan (copia). (Obligatorio)"),
        ("precio_plan", "NUMERIC", "14,0", "Precio del plan (copia). (Obligatorio)"),
        ("cantidad_reportes", "INTEGER", "-", "Reportes incluidos (copia). (Obligatorio)"),
        ("precio_sobre_fact", "NUMERIC", "14,0", "Precio de excedente (copia). (Obligatorio)"),
        ("fecha_desde", "TIMESTAMP", "-", "Inicio de vigencia. (Obligatorio)"),
        ("fecha_hasta", "TIMESTAMP", "-", "Fin de vigencia. (Opcional)"),
        ("cambiado_por", "INTEGER (FK)", "-", "Usuario que realizó el cambio. (ref. usuarios)"),
        ("fecha_cambio", "TIMESTAMP", "-", "Fecha del cambio."),
    ]),
    ("auditoria", "Bitácora centralizada que registra los eventos y acciones realizadas en el sistema.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del evento. (Autoincremental; Obligatorio)"),
        ("usuario_id", "INTEGER", "-", "Usuario que ejecutó la acción. (Opcional)"),
        ("usuario_nombre", "VARCHAR", "150", "Nombre del usuario (copia al momento del evento). (Opcional)"),
        ("id_empresa", "INTEGER", "-", "Empresa asociada al evento. (Opcional)"),
        ("tipo_evento", "VARCHAR", "30", "AUTENTICACION, SEGURIDAD, USUARIOS, CLIENTES, EVALUACIONES, REPORTES, FACTURACION. (Obligatorio)"),
        ("accion", "VARCHAR", "50", "Acción realizada (ej. LOGIN_EXITOSO). (Obligatorio)"),
        ("modulo", "VARCHAR", "50", "Módulo del sistema. (Opcional)"),
        ("entidad", "VARCHAR", "50", "Entidad afectada. (Opcional)"),
        ("registro_id", "VARCHAR", "50", "ID del registro afectado. (Opcional)"),
        ("resultado", "VARCHAR", "20", "EXITO, FALLO. (Por defecto: EXITO)"),
        ("ip", "VARCHAR", "50", "IP de origen. (Opcional)"),
        ("info_adicional", "TEXT", "-", "Información adicional en formato JSON. (Opcional)"),
        ("valores_anteriores", "TEXT", "-", "Valores anteriores en JSON. (Opcional)"),
        ("valores_nuevos", "TEXT", "-", "Valores nuevos en JSON. (Opcional)"),
        ("fecha", "TIMESTAMP", "-", "Fecha y hora del evento."),
    ]),
    ("scoring_modelo", "Versión del modelo de scoring del núcleo (permite versionamiento).", [
        ("id", "INTEGER (PK)", "-", "Identificador único del modelo. (Autoincremental; Obligatorio)"),
        ("nombre", "VARCHAR", "100", "Nombre del modelo. (Obligatorio)"),
        ("version", "VARCHAR", "20", "Versión del modelo. (Por defecto: 1.0)"),
        ("descripcion", "VARCHAR", "500", "Descripción del modelo. (Opcional)"),
        ("id_empresa", "INTEGER", "-", "Empresa dueña (0 = global). (Obligatorio)"),
        ("activo", "BOOLEAN", "-", "Indica si el modelo está activo. (Por defecto: verdadero)"),
        ("creado_en", "TIMESTAMP", "-", "Fecha de creación."),
    ]),
    ("scoring_factor", "Factor de evaluación configurable de un modelo de scoring.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del factor. (Autoincremental; Obligatorio)"),
        ("modelo_id", "INTEGER (FK)", "-", "Modelo al que pertenece. (Obligatorio; ref. scoring_modelo)"),
        ("codigo", "VARCHAR", "50", "Código único del factor dentro del modelo. (Obligatorio)"),
        ("nombre", "VARCHAR", "100", "Nombre del factor. (Obligatorio)"),
        ("descripcion", "VARCHAR", "500", "Descripción del factor. (Opcional)"),
        ("tipo_dato", "VARCHAR", "20", "Tipo de dato del factor. (Por defecto: numerico)"),
        ("tipo_persona", "VARCHAR", "10", "PF, PJ, AMBOS. (Por defecto: AMBOS)"),
        ("categoria", "VARCHAR", "20", "principal, complementario. (Por defecto: principal)"),
        ("obligatorio", "BOOLEAN", "-", "Indica si el factor es obligatorio. (Por defecto: verdadero)"),
        ("activo", "BOOLEAN", "-", "Indica si el factor está activo. (Por defecto: verdadero)"),
        ("orden", "INTEGER", "-", "Orden de evaluación. (Por defecto: 0)"),
    ]),
    ("scoring_catalogo", "Opciones válidas para factores de tipo catálogo.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la opción. (Autoincremental; Obligatorio)"),
        ("factor_id", "INTEGER (FK)", "-", "Factor asociado. (Obligatorio; ref. scoring_factor)"),
        ("valor", "VARCHAR", "100", "Valor interno de la opción. (Obligatorio)"),
        ("etiqueta", "VARCHAR", "100", "Etiqueta visible de la opción. (Obligatorio)"),
        ("orden", "INTEGER", "-", "Orden de presentación. (Por defecto: 0)"),
    ]),
    ("scoring_regla", "Regla IF-THEN de un factor de scoring.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la regla. (Autoincremental; Obligatorio)"),
        ("factor_id", "INTEGER (FK)", "-", "Factor al que pertenece. (Obligatorio; ref. scoring_factor)"),
        ("nombre", "VARCHAR", "100", "Nombre de la regla. (Obligatorio)"),
        ("operador", "VARCHAR", "10", ">=, <=, ==, entre, en. (Obligatorio)"),
        ("valor_min", "VARCHAR", "100", "Valor o límite inferior. (Opcional)"),
        ("valor_max", "VARCHAR", "100", "Límite superior (para 'entre'). (Opcional)"),
        ("nivel", "VARCHAR", "10", "bajo, medio, alto. (Obligatorio)"),
        ("peso", "NUMERIC", "8,2", "Peso (positivo=desfavorable, negativo=favorable). (Obligatorio)"),
        ("explicacion", "VARCHAR", "500", "Texto explicativo. (Opcional)"),
        ("orden", "INTEGER", "-", "Orden de evaluación. (Por defecto: 0)"),
    ]),
    ("scoring_umbral", "Umbrales de clasificación global del score.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del umbral. (Autoincremental; Obligatorio)"),
        ("modelo_id", "INTEGER (FK)", "-", "Modelo al que pertenece. (Obligatorio; ref. scoring_modelo)"),
        ("nivel", "VARCHAR", "10", "bajo, medio, alto. (Obligatorio)"),
        ("score_min", "NUMERIC", "8,2", "Score mínimo del rango. (Obligatorio)"),
        ("score_max", "NUMERIC", "8,2", "Score máximo del rango. (Obligatorio)"),
        ("descripcion", "VARCHAR", "200", "Descripción del umbral. (Opcional)"),
    ]),
    ("scoring_evaluacion", "Resultado completo de una evaluación del núcleo de scoring.", [
        ("id", "INTEGER (PK)", "-", "Identificador único de la evaluación. (Autoincremental; Obligatorio)"),
        ("cliente_id", "INTEGER", "-", "Cliente evaluado. (Obligatorio)"),
        ("modelo_id", "INTEGER (FK)", "-", "Modelo aplicado. (Obligatorio; ref. scoring_modelo)"),
        ("modelo_version", "VARCHAR", "20", "Versión del modelo aplicado. (Obligatorio)"),
        ("usuario_id", "INTEGER", "-", "Usuario que ejecutó la evaluación. (Obligatorio)"),
        ("id_empresa", "INTEGER", "-", "Empresa asociada (0 = global). (Obligatorio)"),
        ("tipo_persona", "VARCHAR", "10", "PF, PJ. (Obligatorio)"),
        ("fecha", "TIMESTAMP", "-", "Fecha de la evaluación."),
        ("score_total", "NUMERIC", "8,2", "Score total obtenido. (Obligatorio)"),
        ("clasificacion", "VARCHAR", "10", "bajo, medio, alto. (Obligatorio)"),
        ("explicacion", "TEXT", "-", "Explicación global del resultado. (Opcional)"),
        ("factores_evaluados", "INTEGER", "-", "Cantidad de factores evaluados. (Por defecto: 0)"),
        ("factores_sin_dato", "INTEGER", "-", "Cantidad de factores sin dato. (Por defecto: 0)"),
        ("estado", "VARCHAR", "20", "completa, incompleta. (Por defecto: completa)"),
    ]),
    ("scoring_detalle", "Resultado individual por factor de una evaluación de scoring.", [
        ("id", "INTEGER (PK)", "-", "Identificador único del detalle. (Autoincremental; Obligatorio)"),
        ("evaluacion_id", "INTEGER (FK)", "-", "Evaluación asociada. (Obligatorio; ref. scoring_evaluacion)"),
        ("factor_codigo", "VARCHAR", "50", "Código del factor evaluado. (Obligatorio)"),
        ("factor_nombre", "VARCHAR", "100", "Nombre del factor. (Obligatorio)"),
        ("categoria", "VARCHAR", "20", "principal, complementario. (Opcional)"),
        ("valor_original", "VARCHAR", "200", "Valor original del cliente. (Opcional)"),
        ("estado", "VARCHAR", "20", "evaluado, sin_dato, no_aplica, invalido. (Obligatorio)"),
        ("regla_aplicada", "VARCHAR", "100", "Nombre de la regla aplicada. (Opcional)"),
        ("nivel", "VARCHAR", "10", "bajo, medio, alto. (Opcional)"),
        ("peso", "NUMERIC", "8,2", "Peso aplicado. (Por defecto: 0)"),
        ("explicacion", "VARCHAR", "500", "Explicación del resultado. (Opcional)"),
    ]),
]


def set_cell_bg(cell, color_hex):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def main():
    doc = Document()

    # Estilo base
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10)

    # Título
    t = doc.add_heading('Diccionario de Datos', level=0)
    p = doc.add_paragraph(
        'Sistema Experto de Evaluación de Clientes — Inmobiliaria. '
        'Motor de base de datos: PostgreSQL. A continuación se detalla, por cada tabla, '
        'la estructura de campos con su tipo de dato, longitud y descripción/restricciones.'
    )
    p.runs[0].italic = True

    for nombre, descripcion, campos in TABLAS:
        doc.add_paragraph()
        h = doc.add_heading(f'Tabla: {nombre}', level=2)
        d = doc.add_paragraph()
        r = d.add_run('Descripción: ')
        r.bold = True
        d.add_run(descripcion)

        tabla = doc.add_table(rows=1, cols=4)
        tabla.style = 'Table Grid'
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Encabezado
        hdr = tabla.rows[0].cells
        titulos = ['Campo', 'Tipo', 'Long.', 'Descripción / Restricciones']
        for i, tit in enumerate(titulos):
            hdr[i].text = ''
            par = hdr[i].paragraphs[0]
            run = par.add_run(tit)
            run.bold = True
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_bg(hdr[i], 'D9D9D9')

        # Filas
        for campo, tipo, longitud, desc in campos:
            fila = tabla.add_row().cells
            fila[0].text = campo
            fila[1].text = tipo
            fila[2].text = longitud
            fila[3].text = desc

        # Ancho de columnas
        from docx.shared import Cm
        anchos = [Cm(4.0), Cm(3.2), Cm(1.6), Cm(8.5)]
        for row in tabla.rows:
            for i, cell in enumerate(row.cells):
                cell.width = anchos[i]

    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    doc.save(SALIDA)
    print("Documento generado:", SALIDA)
    print("Total de tablas:", len(TABLAS))


if __name__ == "__main__":
    main()
